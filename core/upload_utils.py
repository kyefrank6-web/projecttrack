from __future__ import annotations

import re


SUPERVISOR_HEADER_HINTS = {
    "full_name",
    "name",
    "email",
    "username",
    "supervisor",
    "supervisor_name",
    "e_mail",
    "mail",
    "login",
    "user",
    "user_name",
}

STUDENT_HEADER_HINTS = {
    "full_name",
    "name",
    "student_name",
    "class_level",
    "class",
    "form",
    "level",
    "supervisor_username",
    "supervisor",
    "student_no",
    "reg_no",
    "stream",
}

SUPERVISOR_DEFAULT_COLUMNS = ["full_name", "email", "username"]
STUDENT_DEFAULT_COLUMNS = ["full_name", "student_no", "class_level", "stream", "supervisor_username"]
SUPERVISOR_SIMPLE_COLUMNS = ["full_name"]


def normalize_header(key: str) -> str:
    """Turn 'Full Name', 'full name', 'FULL_NAME' into 'full_name'."""
    k = (key or "").strip().lower()
    k = re.sub(r"[^\w\s\-]", "", k)
    k = re.sub(r"[\s\-]+", "_", k)
    return k.strip("_")


def headers_look_like_column_names(header_cells: list[str], hints: set[str]) -> bool:
    norms = {normalize_header(h) for h in header_cells if (h or "").strip()}
    if not norms:
        return False
    if norms & hints:
        return True
    return any(
        token in n
        for n in norms
        for token in ("name", "email", "class", "supervisor", "stream", "username", "level", "form")
    )


def map_row_from_cells(cells: list[str], column_names: list[str]) -> dict[str, str]:
    row: dict[str, str] = {}
    for i, col in enumerate(column_names):
        val = cells[i] if i < len(cells) else ""
        row[col] = ("" if val is None else str(val)).strip()
    return row


def parse_raw_table(
    raw_rows: list[list],
    *,
    default_columns: list[str] | None,
    header_hints: set[str],
) -> list[dict[str, str]]:
    if not raw_rows:
        return []

    first_cells = [("" if c is None else str(c).strip()) for c in raw_rows[0]]
    use_defaults = bool(default_columns) and not headers_look_like_column_names(first_cells, header_hints)

    if use_defaults:
        columns = default_columns
        data_rows = raw_rows
    else:
        columns = [h if h else f"column_{i + 1}" for i, h in enumerate(first_cells)]
        data_rows = raw_rows[1:]

    out: list[dict[str, str]] = []
    for r in data_rows:
        cells = [("" if c is None else str(c).strip()) for c in r]
        if not any(cells):
            continue
        out.append(map_row_from_cells(cells, columns))
    return out


def row_lookup(row: dict[str, str], *aliases: str, allow_lone_value: bool = False) -> str:
    """Find a cell value using flexible column header names.

    When allow_lone_value is True and the row has a single non-empty cell,
    return that cell (used for name-only uploads). Never enable this for
    email/username lookups — otherwise the person's name is mistaken for a login.
    """
    norm = {normalize_header(k): (v or "").strip() for k, v in row.items()}
    for alias in aliases:
        val = norm.get(normalize_header(alias), "")
        if val:
            return val
    if allow_lone_value:
        non_empty = [v for v in row.values() if (v or "").strip()]
        if len(non_empty) == 1:
            return non_empty[0]
    return ""


def headers_found(row: dict[str, str]) -> str:
    return ", ".join(sorted({normalize_header(k) for k in row.keys() if k})) or "(none)"


def read_supervisor_upload_table(uploaded_file) -> list[dict[str, str]]:
    """
    Parse a supervisor upload file.

    Supports:
    - A single column of names (no header row required)
    - A header row with name / optional email / optional username columns
    """
    raw_rows = load_upload_rows(uploaded_file)
    if not raw_rows:
        return []

    first_cells = [("" if c is None else str(c).strip()) for c in raw_rows[0]]
    if headers_look_like_column_names(first_cells, SUPERVISOR_HEADER_HINTS):
        return parse_raw_table(
            raw_rows,
            default_columns=SUPERVISOR_DEFAULT_COLUMNS,
            header_hints=SUPERVISOR_HEADER_HINTS,
        )

    non_empty_counts = [
        sum(1 for c in (("" if cell is None else str(cell).strip()) for cell in row) if c)
        for row in raw_rows[:10]
    ]
    max_cols = max(non_empty_counts, default=0)
    default_columns = SUPERVISOR_DEFAULT_COLUMNS if max_cols >= 2 else SUPERVISOR_SIMPLE_COLUMNS
    return parse_raw_table(
        raw_rows,
        default_columns=default_columns,
        header_hints=SUPERVISOR_HEADER_HINTS,
    )


def load_upload_rows(uploaded_file) -> list[list]:
    import csv
    import io

    from openpyxl import load_workbook

    # Form / previous reads may leave the pointer at EOF.
    if hasattr(uploaded_file, "seek"):
        try:
            uploaded_file.seek(0)
        except Exception:
            pass

    name = (getattr(uploaded_file, "name", "") or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ValueError(
            "Old Excel .xls files are not supported. "
            "Save the file as .xlsx (Excel Workbook) or CSV and try again."
        )

    raw = uploaded_file.read()
    if not raw:
        return []

    if name.endswith(".xlsx"):
        try:
            # BytesIO is more reliable than UploadedFile handles on Render/Gunicorn.
            wb = load_workbook(io.BytesIO(raw), data_only=True)
            ws = wb.active
            return [list(r) for r in ws.iter_rows(values_only=True)]
        except Exception as exc:
            raise ValueError(
                "Could not read that Excel file. "
                "Use a .xlsx workbook or CSV (UTF-8), and make sure the file is not corrupted."
            ) from exc

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return list(csv.reader(io.StringIO(text)))
