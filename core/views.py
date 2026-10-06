from __future__ import annotations

import csv
import io
import json
import re
import secrets
import zipfile
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.conf import settings
from django.core.files import File
from django.contrib.auth import get_user_model
from django.contrib.auth import authenticate, login, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib.auth.hashers import make_password
from django.contrib.auth.views import LoginView as DjangoLoginView
from django.contrib.auth.views import LogoutView as DjangoLogoutView
from django.db import transaction
from django.db.models import Count, Q
from django.db.models.functions import Lower
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme

from openpyxl import Workbook

from .forms import (
    AcademicPromotionForm,
    ClassFilterForm,
    ClassProjectThemeForm,
    CohortProjectThemeForm,
    ManualSupervisorForm,
    OverallStudentForm,
    ObservationChecklistUploadForm,
    PasswordResetRequestForm,
    RegisterSchoolForm,
    SchoolBrandingForm,
    SchoolPasswordChangeForm,
    SchoolRejectForm,
    SchoolDeleteConfirmForm,
    ScoreStudentForm,
    SetNewPasswordForm,
    StudentEditForm,
    UploadFileForm,
    UsernameChangeForm,
)
from .project_theme_utils import cohort_year_for_class, get_student_project_theme, project_themes_by_class
from .builtin_checklists import (
    BUILTIN_CHECKLISTS,
    checklist_document_for,
    checklist_headings_for,
    cohort_upload_guide,
    builtin_format_years,
    get_builtin_for_cohort,
    get_builtin_format,
)
from .observation_defaults import OBSERVATION_RATING_LABELS_CHECKBOX, OBSERVATION_RATING_LABELS_SCALE
from .observation_import import import_structure_file
from .observation_utils import (
    existing_ratings_map,
    get_observation_structure,
    latest_assessor_display_name,
    student_has_observation_ratings,
    active_checklist_class_levels,
    student_display_score_maps,
    terms_with_marks_for_students,
    percentages_to_persist,
    resolve_checklist,
    resolve_checklist_for_student,
    get_submitted_competency_ids,
    save_competency_observation_ratings_and_score,
    save_observation_ratings_and_scores,
    seed_checklist_criteria,
    submitted_competency_ids_for_assessment,
    uses_checkbox_scoring,
)
from .profile_utils import get_active_profile
from .themes import THEMES
from .school_utils import (
    SCHOOL_DELETION_GRACE_DAYS,
    clear_registration_blockers,
    current_school_term,
    deactivate_school,
    delete_all_students_from_school,
    delete_school_completely,
    delete_student_completely,
    permanent_deletion_at,
    purge_expired_deactivated_schools,
    ensure_supervisor_membership,
    remove_all_supervisors_from_school,
    remove_supervisor_from_school,
    restore_deactivated_school,
)
from .upload_utils import (
    SUPERVISOR_DEFAULT_COLUMNS,
    SUPERVISOR_HEADER_HINTS,
    SUPERVISOR_SIMPLE_COLUMNS,
    STUDENT_DEFAULT_COLUMNS,
    STUDENT_HEADER_HINTS,
    headers_found,
    load_upload_rows,
    parse_raw_table,
    read_supervisor_upload_table,
    row_lookup,
)
from .evidence_utils import evidence_complete, evidence_completion_count, student_evidence_slots
from .pdf_reports import (
    build_class_evidence_pdf,
    build_class_project_titles_pdf,
    build_class_scores_pdf,
    build_student_scored_observation_pdf,
)
from .promotion_utils import (
    PromotionAlreadyRunError,
    latest_promotion_run,
    preview_academic_promotion,
    run_academic_promotion,
)
from .models import (
    AcademicPromotionRun,
    AssessmentScheme,
    Competency,
    CompetencyScore,
    EvidenceCategory,
    ProjectAssessment,
    ClassProjectTheme,
    CohortProjectTheme,
    ObservationChecklist,
    ProjectEvidence,
    Role,
    School,
    SchoolRegistrationStatus,
    SecondaryClassLevel,
    Student,
    UserProfile,
)


User = get_user_model()


class LoginView(DjangoLoginView):
    template_name = "auth/login.html"

    def post(self, request, *args, **kwargs):
        username = (request.POST.get("username") or "").strip()
        password = request.POST.get("password") or ""

        user = authenticate(request, username=username, password=password)
        if user is None:
            from .maintenance_utils import is_maintenance_mode, maintenance_message

            if is_maintenance_mode():
                messages.error(request, maintenance_message())
            else:
                messages.error(
                    request,
                    "Invalid username or password. Supervisors can use their account password "
                    "even when registered at more than one school.",
                )
            return self.render_to_response(self.get_context_data())

        login(request, user)
        profile = get_active_profile(request)
        if profile and profile.role != Role.SUPERADMIN and profile.school_id:
            school = profile.school
            if school.status == SchoolRegistrationStatus.PENDING:
                messages.warning(
                    request,
                    "Your school registration is still pending approval by the platform administrator.",
                )
            elif school.status == SchoolRegistrationStatus.REJECTED:
                messages.error(
                    request,
                    "Your school registration was rejected. "
                    + (school.rejection_reason or "Contact the administrator for details."),
                )
                return self.render_to_response(self.get_context_data())
            elif school.status == SchoolRegistrationStatus.DEACTIVATED:
                messages.error(
                    request,
                    "Your school account has been deactivated and is scheduled for permanent removal.",
                )
                return self.render_to_response(self.get_context_data())

        if profile and profile.must_change_password:
            messages.info(
                request,
                "You signed in with a one-time password. Please set a new password to continue.",
            )
            return redirect("password_change")

        return redirect(self.get_success_url())

    def get_success_url(self):
        return self.get_redirect_url() or str(self.success_url or "/")


class LogoutView(DjangoLogoutView):
    next_page = "home"


def password_reset_request(request):
    """Self-service reset via email (username + email + optional school name)."""
    from .password_reset_utils import membership_label, send_password_reset_email

    if request.user.is_authenticated:
        return redirect("home")

    form = PasswordResetRequestForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        membership = form.cleaned_data["membership"]
        email = form.cleaned_data["email"]
        try:
            send_password_reset_email(
                request=request,
                membership=membership,
                recipient_email=email,
            )
        except Exception:
            messages.error(
                request,
                "Could not send the reset email. Ask your overall supervisor to reset your password, "
                "or contact the platform administrator.",
            )
        else:
            messages.success(
                request,
                f"If the details were correct, a reset link was sent to {email} for "
                f"{membership_label(membership)}. Check your inbox (and spam folder).",
            )
        return redirect("password_reset_done")

    return render(request, "auth/password_reset_request.html", {"form": form})


def password_reset_done(request):
    return render(request, "auth/password_reset_done.html")


def password_reset_confirm(request, signed_membership_id: str, token: str):
    from .password_reset_utils import (
        SchoolPasswordResetTokenGenerator,
        membership_label,
        set_membership_password,
        unsign_membership_id,
    )

    membership_id = unsign_membership_id(signed_membership_id)
    if not membership_id:
        messages.error(request, "This reset link is invalid or has expired.")
        return redirect("password_reset_request")

    membership = (
        UserProfile.objects.filter(pk=membership_id).select_related("user", "school").first()
    )
    if not membership:
        messages.error(request, "This reset link is invalid.")
        return redirect("password_reset_request")

    token_ok = SchoolPasswordResetTokenGenerator(membership.id).check_token(
        membership.user, token
    )
    if not token_ok:
        messages.error(request, "This reset link is invalid or has already been used.")
        return redirect("password_reset_request")

    form = SetNewPasswordForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        set_membership_password(
            membership, form.cleaned_data["new_password1"], must_change=False
        )
        messages.success(
            request,
            f"Password updated for {membership_label(membership)}. You can sign in now.",
        )
        return redirect("password_reset_complete")

    return render(
        request,
        "auth/password_reset_confirm.html",
        {"form": form, "membership": membership},
    )


def password_reset_complete(request):
    return render(request, "auth/password_reset_complete.html")


def _require_profile(request) -> UserProfile:
    profile = get_active_profile(request)
    if not profile:
        raise Http404("Profile not found")
    return profile


def _is_superadmin(profile: UserProfile) -> bool:
    return profile.role == Role.SUPERADMIN


def home(request):
    """Landing page for everyone; after login users arrive here first."""
    from django.templatetags.static import static as static_url

    return render(
        request,
        "welcome.html",
        {
            "welcome_photo_url": static_url("core/img/student-project-work.jpg"),
        },
    )


def register_school(request):
    """
    Public page to create a new School and its first Overall Supervisor account.
    """

    if request.user.is_authenticated:
        return redirect("home")

    if request.method == "POST":
        form = RegisterSchoolForm(request.POST)
        if form.is_valid():
            school_name = form.cleaned_data["school_name"].strip()
            username = form.cleaned_data["overall_username"].strip()
            email = form.cleaned_data.get("overall_email") or ""
            password = form.cleaned_data["overall_password"]

            with transaction.atomic():
                clear_registration_blockers(school_name=school_name, username=username)
                user = User.objects.create_user(username=username, email=email, password=password)
                user.first_name = school_name
                user.save(update_fields=["first_name"])

                school = School.objects.create(
                    name=school_name,
                    overall_supervisor=user,
                    status=SchoolRegistrationStatus.PENDING,
                )
                UserProfile.objects.create(
                    user=user,
                    school=school,
                    role=Role.OVERALL_SUPERVISOR,
                    school_password=user.password,
                )

                # Seed a default scheme if none exists.
                if not AssessmentScheme.objects.exists():
                    scheme = AssessmentScheme.objects.create(name="Uganda Secondary Project (Default)", active=True)
                    for i, name in enumerate(
                        [
                            "Project planning",
                            "Project implementation",
                            "Project reporting",
                            "Project dissemination",
                        ],
                        start=1,
                    ):
                        Competency.objects.create(scheme=scheme, name=name, order=i)

            messages.success(
                request,
                "Registration submitted. A platform administrator will review your school. "
                "You can log in after approval. If your school was deleted or rejected before, "
                "register again here and wait for approval.",
            )
            return redirect("login")
    else:
        form = RegisterSchoolForm()

    return render(request, "auth/register_school.html", {"form": form})


def _is_overall(profile: UserProfile) -> bool:
    return profile.role == Role.OVERALL_SUPERVISOR


def _is_supervisor(profile: UserProfile) -> bool:
    return profile.role == Role.SUPERVISOR


def _user_display_name(user) -> str:
    full = (user.get_full_name() or "").strip()
    return full or user.username


def _class_level_sort_key(code: str) -> int:
    order = {c: i for i, (c, _) in enumerate(SecondaryClassLevel.choices)}
    return order.get(code, 99)


def _supervisor_student_classes(supervisor_user, school: School) -> list[str]:
    codes = (
        Student.objects.active_only()
        .filter(school=school, supervisor=supervisor_user)
        .values_list("class_level", flat=True)
        .distinct()
    )
    return sorted(set(codes), key=_class_level_sort_key)


def _resolve_supervisor_class_filter(request, assigned_classes: list[str]) -> str:
    selected = (request.GET.get("class") or "").strip().upper()
    if selected in assigned_classes:
        return selected
    return assigned_classes[0] if assigned_classes else ""


def _supervisor_profiles_queryset(school: School, query: str = ""):
    qs = (
        UserProfile.objects.filter(school=school, role=Role.SUPERVISOR)
        .select_related("user")
        .annotate(
            student_count=Count(
                "user__supervised_students",
                filter=Q(
                    user__supervised_students__school_id=school.id,
                    user__supervised_students__active=True,
                ),
            )
        )
        .order_by(Lower("user__first_name"), Lower("user__username"))
    )
    q = (query or "").strip()
    if q:
        qs = qs.filter(
            Q(user__username__icontains=q)
            | Q(user__first_name__icontains=q)
            | Q(user__last_name__icontains=q)
            | Q(user__email__icontains=q)
        )
    return qs


def _overall_students_queryset(
    school: School,
    query: str = "",
    *,
    supervisor_user_id: int | None = None,
    unassigned_only: bool = False,
):
    qs = Student.objects.active_only().filter(school=school).select_related("supervisor").order_by(
        Lower("full_name"), "class_level"
    )
    if unassigned_only:
        qs = qs.filter(supervisor__isnull=True)
    elif supervisor_user_id is not None:
        qs = qs.filter(supervisor_id=supervisor_user_id)
    q = (query or "").strip()
    if q:
        qs = qs.filter(
            Q(full_name__icontains=q)
            | Q(student_no__icontains=q)
            | Q(project_title__icontains=q)
            | Q(class_level__icontains=q)
            | Q(stream__icontains=q)
            | Q(supervisor__username__icontains=q)
            | Q(supervisor__first_name__icontains=q)
            | Q(supervisor__last_name__icontains=q)
        )
    return qs


def _unassigned_student_count(school: School) -> int:
    return Student.objects.active_only().filter(school=school, supervisor__isnull=True).count()


def _school_supervisor_profiles(school: School):
    return UserProfile.objects.filter(school=school, role=Role.SUPERVISOR).select_related(
        "user"
    ).order_by(Lower("user__first_name"), Lower("user__username"))


def _redirect_preserve_get(request, url_name: str, **url_kwargs):
    referer = request.META.get("HTTP_REFERER", "")
    if referer and url_has_allowed_host_and_scheme(
        referer, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        return redirect(referer)
    url = reverse(url_name, kwargs=url_kwargs)
    qs = request.GET.urlencode()
    if qs:
        url = f"{url}?{qs}"
    return redirect(url)


def _require_school_supervisor(profile: UserProfile, supervisor_user_id: int):
    membership = UserProfile.objects.filter(
        school=profile.school,
        role=Role.SUPERVISOR,
        user_id=supervisor_user_id,
    ).select_related("user").first()
    if not membership:
        raise Http404()
    return membership


def _unique_username(
    base: str,
    *,
    exclude_user_id: int | None = None,
    reserved: set[str] | None = None,
) -> str:
    candidate = base or "user"
    suffix = 1
    reserved = reserved or set()
    qs = User.objects.all()
    if exclude_user_id:
        qs = qs.exclude(pk=exclude_user_id)
    while qs.filter(username=candidate).exists() or candidate in reserved:
        suffix += 1
        candidate = f"{base}{suffix}"
    return candidate


def _username_base_from_identity(*, full_name: str = "", email: str = "", username: str = "") -> str:
    if username:
        cleaned = "".join(ch for ch in username.strip().lower() if ch.isalnum() or ch in "._-")
        return (cleaned[:20] or "user")
    if email and "@" in email:
        base = email.split("@")[0].lower()
        cleaned = "".join(ch for ch in base if ch.isalnum() or ch in "._-")[:20]
        return cleaned or "user"

    titles = {"mr", "mrs", "ms", "miss", "dr", "sir", "madam", "prof", "rev"}
    parts = [
        "".join(ch for ch in part.lower() if ch.isalnum())
        for part in (full_name or "").replace(".", " ").split()
    ]
    parts = [p for p in parts if p and p not in titles]
    if not parts:
        return "user"
    # Prefer surname-like token when available, else first given name.
    base = parts[-1] if len(parts) > 1 else parts[0]
    return base[:20] or "user"


def _provision_school_supervisor(
    *,
    school: School,
    full_name: str,
    email: str = "",
    username: str = "",
    password: str | None = None,
) -> dict[str, str]:
    """
    Create a new supervisor account or attach an existing username to this school.
    Same display names are allowed. Auto-generated usernames are made unique so
    people with the same name get separate accounts. Explicit usernames still link
    an existing login across schools when intended.
    """
    explicit_username = bool((username or "").strip())
    if not password:
        password = secrets.token_urlsafe(10)
    password_hash = make_password(password)

    if explicit_username:
        username = username.strip()
        existing_user = User.objects.filter(username=username).first()
        if existing_user:
            membership, created = ensure_supervisor_membership(
                school=school,
                user=existing_user,
                password_hash=password_hash,
                must_change_password=True,
            )
            existing_user.first_name = full_name
            if email:
                existing_user.email = email
                existing_user.save(update_fields=["first_name", "email"])
            else:
                existing_user.save(update_fields=["first_name"])
            if not created and school.overall_supervisor_id == existing_user.id:
                note = (
                    "Overall supervisor also added as supervisor — use the supervisor "
                    "password shown here to sign in with the supervisor role."
                )
            elif not created:
                note = (
                    "Existing account linked to this school — use this school's password when logging in."
                )
            else:
                note = "Existing username added to this school."
            return {
                "full_name": full_name,
                "username": username,
                "email": email or existing_user.email,
                "password": password,
                "school": school.name,
                "note": note,
            }
    else:
        username = _unique_username(
            _username_base_from_identity(full_name=full_name, email=email)
        )

    user = User.objects.create_user(username=username, email=email, password=password)
    user.first_name = full_name
    user.save(update_fields=["first_name"])
    UserProfile.objects.create(
        user=user,
        school=school,
        role=Role.SUPERVISOR,
        school_password=user.password,
        must_change_password=True,
    )
    return {
        "full_name": full_name,
        "username": username,
        "email": email,
        "password": password,
        "school": school.name,
        "note": "",
    }


@login_required
def overall_dashboard(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    school = profile.school
    search_q = (request.GET.get("q") or "").strip()
    supervisors = list(_supervisor_profiles_queryset(school, search_q))
    supervisors_count = UserProfile.objects.filter(school=school, role=Role.SUPERVISOR).count()
    students_count = Student.objects.active_only().filter(school=school).count()
    search_students = (
        list(_overall_students_queryset(school, search_q)[:50]) if search_q else []
    )

    return render(
        request,
        "overall/dashboard.html",
        {
            "school": school,
            "supervisors_count": supervisors_count,
            "students_count": students_count,
            "class_levels": [c for c, _ in SecondaryClassLevel.choices],
            "supervisors": supervisors,
            "search_q": search_q,
            "search_students": search_students,
        },
    )


@login_required
def overall_supervisors(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    search_q = (request.GET.get("q") or "").strip()
    supervisors = list(_supervisor_profiles_queryset(profile.school, search_q))

    return render(
        request,
        "overall/supervisors.html",
        {
            "supervisors": supervisors,
            "search_q": search_q,
            "supervisors_count": UserProfile.objects.filter(
                school=profile.school, role=Role.SUPERVISOR
            ).count(),
        },
    )


@login_required
def overall_delete_all_supervisors(request):
    """Remove every regular supervisor from this school (not the overall supervisor)."""
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()
    if request.method != "POST":
        raise Http404()

    removed, unassigned = remove_all_supervisors_from_school(school=profile.school)
    if removed:
        msg = f"Removed {removed} supervisor(s) from this school."
        if unassigned:
            msg += f" {unassigned} learner(s) are now without a supervisor."
        messages.success(request, msg)
    else:
        messages.info(request, "There are no supervisors to remove.")

    return redirect("overall_supervisors")


@login_required
def overall_students(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    search_q = (request.GET.get("q") or "").strip()
    class_level = (request.GET.get("class") or "").strip().upper()
    valid_classes = {c for c, _ in SecondaryClassLevel.choices}

    students_qs = _overall_students_queryset(profile.school, search_q)
    if class_level in valid_classes:
        students_qs = students_qs.filter(class_level=class_level)

    students = list(students_qs)
    total_students = Student.objects.active_only().filter(school=profile.school).count()

    class_counts = {
        row["class_level"]: row["count"]
        for row in Student.objects.active_only()
        .filter(school=profile.school)
        .values("class_level")
        .annotate(count=Count("id"))
    }

    unassigned_count = _unassigned_student_count(profile.school)

    return render(
        request,
        "overall/students.html",
        {
            "students": students,
            "search_q": search_q,
            "class_level": class_level if class_level in valid_classes else "",
            "class_levels": [c for c, _ in SecondaryClassLevel.choices],
            "total_students": total_students,
            "filtered_students": len(students),
            "class_counts": class_counts,
            "unassigned_count": unassigned_count,
            "unassigned_only": False,
        },
    )


@login_required
def overall_students_unassigned(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    search_q = (request.GET.get("q") or "").strip()
    class_level = (request.GET.get("class") or "").strip().upper()
    valid_classes = {c for c, _ in SecondaryClassLevel.choices}

    students_qs = _overall_students_queryset(
        profile.school, search_q, unassigned_only=True
    )
    if class_level in valid_classes:
        students_qs = students_qs.filter(class_level=class_level)

    students = list(students_qs)
    unassigned_count = _unassigned_student_count(profile.school)
    total_students = Student.objects.active_only().filter(school=profile.school).count()

    school_supervisors = list(_school_supervisor_profiles(profile.school))

    class_counts = {
        row["class_level"]: row["count"]
        for row in Student.objects.active_only()
        .filter(school=profile.school, supervisor__isnull=True)
        .values("class_level")
        .annotate(count=Count("id"))
    }

    return render(
        request,
        "overall/students.html",
        {
            "students": students,
            "search_q": search_q,
            "class_level": class_level if class_level in valid_classes else "",
            "class_levels": [c for c, _ in SecondaryClassLevel.choices],
            "total_students": total_students,
            "filtered_students": len(students),
            "class_counts": class_counts,
            "unassigned_count": unassigned_count,
            "unassigned_only": True,
            "school_supervisors": school_supervisors,
        },
    )


@login_required
def overall_promote_students(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    school = profile.school
    default_year = timezone.now().year

    if request.method == "POST":
        form = AcademicPromotionForm(request.POST)
        if form.is_valid():
            from_year = form.cleaned_data["from_academic_year"]
            try:
                run = run_academic_promotion(
                    school.id,
                    from_academic_year=from_year,
                    run_by_id=request.user.id,
                )
                summary = run.summary or {}
                o_grad = summary.get("S4_O_level_graduated", 0)
                a_grad = summary.get("S6_A_level_graduated", 0)
                messages.success(
                    request,
                    f"End-of-year promotion complete for {from_year} → {run.to_academic_year}: "
                    f"{run.promoted_count} learner(s) moved up; "
                    f"{o_grad} O-level S.4 leaver(s) and {a_grad} A-level S.6 leaver(s) graduated. "
                    f"Supervisors, scores, and cohort themes are unchanged. "
                    f"Upload new S.5 A-level learners separately when ready.",
                )
            except PromotionAlreadyRunError:
                messages.error(
                    request,
                    f"Promotion for academic year {from_year} was already completed for this school.",
                )
            return redirect("overall_promote_students")
        from_year = form.data.get("from_academic_year") or default_year
    else:
        from_year = request.GET.get("from_year") or default_year
        form = AcademicPromotionForm(initial={"from_academic_year": from_year})

    try:
        from_year = int(from_year)
    except (TypeError, ValueError):
        from_year = default_year

    preview = preview_academic_promotion(school.id, from_year)
    promotion_history = list(
        AcademicPromotionRun.objects.filter(school=school).select_related("run_by")[:8]
    )

    return render(
        request,
        "overall/promote_students.html",
        {
            "form": form,
            "preview": preview,
            "promotion_history": promotion_history,
            "last_promotion": latest_promotion_run(school.id),
        },
    )


@login_required
def overall_assign_student_supervisor(request, student_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()
    if request.method != "POST":
        raise Http404()

    student = get_object_or_404(
        Student.objects.active_only().select_related("supervisor"),
        id=student_id,
        school=profile.school,
        supervisor__isnull=True,
    )
    supervisor_id = request.POST.get("supervisor")
    if not supervisor_id or not str(supervisor_id).isdigit():
        messages.error(request, "Select a supervisor from the list.")
        return _redirect_preserve_get(request, "overall_students_unassigned")

    membership = _school_supervisor_profiles(profile.school).filter(
        user_id=int(supervisor_id)
    ).first()
    if not membership:
        messages.error(request, "That supervisor is not registered at your school.")
        return _redirect_preserve_get(request, "overall_students_unassigned")

    student.supervisor = membership.user
    student.save(update_fields=["supervisor"])
    display = _user_display_name(membership.user)
    messages.success(
        request,
        f'Assigned {student.full_name} to {display} (@{membership.user.username}).',
    )
    return _redirect_preserve_get(request, "overall_students_unassigned")


@login_required
def overall_delete_student(request, student_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()
    if request.method != "POST":
        raise Http404()

    student = get_object_or_404(Student, id=student_id, school=profile.school)
    name = student.full_name
    delete_student_completely(student)
    messages.success(request, f'Learner "{name}" has been permanently deleted.')

    referer = request.META.get("HTTP_REFERER", "")
    if referer and url_has_allowed_host_and_scheme(
        referer, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        return redirect(referer)
    redirect_qs = request.GET.urlencode()
    url = reverse("overall_students")
    if redirect_qs:
        url = f"{url}?{redirect_qs}"
    return redirect(url)


@login_required
def overall_delete_all_students(request):
    """Permanently remove every student at this school."""
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()
    if request.method != "POST":
        raise Http404()

    deleted = delete_all_students_from_school(school=profile.school)
    if deleted:
        messages.success(
            request,
            f"Permanently deleted {deleted} student(s), including their scores and evidence.",
        )
    else:
        messages.info(request, "There are no students to delete.")

    return redirect("overall_students")


@login_required
def overall_delete_supervisor(request, supervisor_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()
    if request.method != "POST":
        raise Http404()

    try:
        display, unassigned = remove_supervisor_from_school(
            school=profile.school, supervisor_user_id=supervisor_id
        )
    except ValueError as exc:
        messages.error(request, str(exc))
    else:
        msg = f'Supervisor "{display}" has been removed from this school.'
        if unassigned:
            msg += f" {unassigned} learner(s) are now without a supervisor."
        messages.success(request, msg)

    redirect_qs = request.GET.urlencode()
    url = reverse("overall_supervisors")
    if redirect_qs:
        url = f"{url}?{redirect_qs}"
    return redirect(url)


@login_required
def overall_reset_supervisor_password(request, supervisor_id: int):
    """Overall supervisor issues a new temporary password for a school supervisor."""
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()
    if request.method != "POST":
        raise Http404()

    membership = _require_school_supervisor(profile, supervisor_id)
    new_password = secrets.token_urlsafe(10)
    from .password_reset_utils import set_membership_password

    set_membership_password(membership, new_password, must_change=True)
    display = membership.user.get_full_name() or membership.user.username
    request.session["supervisor_password_reset"] = {
        "username": membership.user.username,
        "password": new_password,
        "display_name": display,
    }
    return redirect("overall_supervisor_password_reset_done")


@login_required
def overall_supervisor_password_reset_done(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    payload = request.session.pop("supervisor_password_reset", None)
    if not payload:
        return redirect("overall_supervisors")

    return render(request, "overall/supervisor_password_reset_done.html", {"reset": payload})


@login_required
def overall_supervisor_students(request, supervisor_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    membership = _require_school_supervisor(profile, supervisor_id)
    supervisor_user = membership.user
    search_q = (request.GET.get("q") or "").strip()
    students = list(
        _overall_students_queryset(
            profile.school, search_q, supervisor_user_id=supervisor_id
        )
    )

    return render(
        request,
        "overall/supervisor_students.html",
        {
            "supervisor": supervisor_user,
            "supervisor_display": _user_display_name(supervisor_user),
            "students": students,
            "search_q": search_q,
            "student_count": Student.objects.active_only().filter(
                school=profile.school, supervisor_id=supervisor_id
            ).count(),
        },
    )


@login_required
def export_supervisor_students(request, supervisor_id: int):
    """Excel list of learners assigned to one supervisor at the overall supervisor's school."""
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    membership = _require_school_supervisor(profile, supervisor_id)
    supervisor_user = membership.user
    display = _user_display_name(supervisor_user)
    students = list(
        _overall_students_queryset(profile.school, supervisor_user_id=supervisor_id)
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Learners"
    ws.append(["Full name", "Student number", "Class", "Stream", "Project title", "Supervisor"])
    for student in students:
        ws.append(
            [
                student.full_name,
                student.student_no,
                student.class_level,
                student.stream,
                student.project_title,
                display,
            ]
        )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    safe_supervisor = re.sub(r"[^A-Za-z0-9._-]+", "_", display).strip("._") or "supervisor"
    safe_school = re.sub(r"[^A-Za-z0-9._-]+", "_", profile.school.name).strip("._") or "school"
    filename = f"{safe_school}_{safe_supervisor}_learners.xlsx"
    resp = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _read_uploaded_table(uploaded_file, *, upload_kind: str = "student") -> list[dict[str, str]]:
    """
    Reads CSV or XLSX. If row 1 is not column headers (e.g. only names), uses default columns.
    upload_kind: 'supervisor' | 'student' | 'supervisor_simple' (name-only list)
    """

    if upload_kind == "supervisor":
        default_columns = SUPERVISOR_DEFAULT_COLUMNS
        header_hints = SUPERVISOR_HEADER_HINTS
    elif upload_kind == "supervisor_simple":
        default_columns = SUPERVISOR_SIMPLE_COLUMNS
        header_hints = SUPERVISOR_HEADER_HINTS
    else:
        default_columns = STUDENT_DEFAULT_COLUMNS
        header_hints = STUDENT_HEADER_HINTS

    raw_rows = load_upload_rows(uploaded_file)
    return parse_raw_table(raw_rows, default_columns=default_columns, header_hints=header_hints)


def _next_unique_username(base: str, reserved: set[str]) -> str:
    """Pick a unique username using an in-memory reserved set (avoids N+1 DB hits)."""
    candidate = (base or "user")[:140]
    if candidate not in reserved:
        reserved.add(candidate)
        return candidate
    suffix = 2
    while True:
        candidate = f"{(base or 'user')[:130]}{suffix}"
        if candidate not in reserved:
            reserved.add(candidate)
            return candidate
        suffix += 1


@login_required
def upload_supervisors(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    if request.method == "POST":
        form = UploadFileForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                rows = read_supervisor_upload_table(form.cleaned_data["file"])
            except ValueError as exc:
                messages.error(request, str(exc))
                return redirect("upload_supervisors")
            except Exception:
                messages.error(
                    request,
                    "Could not read that file. Upload a CSV or Excel (.xlsx) file and try again.",
                )
                return redirect("upload_supervisors")

            if not rows:
                messages.error(request, "File has no rows.")
                return redirect("upload_supervisors")

            if len(rows) > 300:
                messages.error(
                    request,
                    f"This file has {len(rows)} rows. Upload at most 300 supervisors at a time "
                    "(split the file and upload in batches).",
                )
                return redirect("upload_supervisors")

            created_credentials: list[dict[str, str]] = []
            reserved_usernames: set[str] = set(User.objects.values_list("username", flat=True))
            try:
                with transaction.atomic():
                    for idx, row in enumerate(rows, start=1):
                        full_name = row_lookup(
                            row,
                            "full_name",
                            "name",
                            "supervisor_name",
                            "supervisor",
                            "names",
                            allow_lone_value=True,
                        )[:150]
                        email = row_lookup(row, "email", "e_mail", "mail")
                        username = row_lookup(row, "username", "user_name", "login", "user")

                        if not full_name:
                            raise ValueError(
                                f"Row {idx}: missing supervisor name. "
                                f"Put names in column full_name or use a single column of names only. "
                                f"Columns detected: {headers_found(row)}"
                            )

                        password = secrets.token_urlsafe(10)
                        password_hash = make_password(password)

                        if email and "@" not in email:
                            email = ""

                        explicit_username = bool(username)
                        if explicit_username:
                            # Keep only characters Django usernames allow (no spaces).
                            username = _username_base_from_identity(username=username.strip())
                            existing_user = User.objects.filter(username=username).first()
                            if existing_user:
                                _, created = ensure_supervisor_membership(
                                    school=profile.school,
                                    user=existing_user,
                                    password_hash=password_hash,
                                    must_change_password=True,
                                )
                                existing_user.first_name = full_name
                                if email:
                                    existing_user.email = email
                                existing_user.save(
                                    update_fields=["first_name", "email"] if email else ["first_name"]
                                )
                                created_credentials.append(
                                    {
                                        "full_name": full_name,
                                        "username": username,
                                        "email": email or existing_user.email,
                                        "password": password,
                                        "school": profile.school.name,
                                        "note": "Same username at multiple schools — use this school's password when logging in.",
                                    }
                                )
                                reserved_usernames.add(username)
                                continue

                            username = _next_unique_username(username, reserved_usernames)
                        else:
                            # Same display names are allowed — always create a distinct login.
                            username = _next_unique_username(
                                _username_base_from_identity(full_name=full_name, email=email),
                                reserved_usernames,
                            )

                        user = User(
                            username=username,
                            email=email or "",
                            first_name=full_name,
                        )
                        user.password = password_hash
                        user.save()

                        UserProfile.objects.create(
                            user=user,
                            school=profile.school,
                            role=Role.SUPERVISOR,
                            school_password=password_hash,
                            must_change_password=True,
                        )
                        created_credentials.append(
                            {
                                "full_name": full_name,
                                "username": username,
                                "email": email,
                                "password": password,
                                "school": profile.school.name,
                                "note": "",
                            }
                        )
            except ValueError as exc:
                messages.error(request, str(exc))
                return redirect("upload_supervisors")
            except Exception as exc:
                from django.db import IntegrityError

                if isinstance(exc, IntegrityError) and "username" in str(exc).lower():
                    messages.error(
                        request,
                        "One or more usernames already exist. Remove duplicate rows, use unique usernames, "
                        "or leave the username column empty so the system generates unique logins.",
                    )
                    return redirect("upload_supervisors")
                messages.error(
                    request,
                    "Upload failed unexpectedly. Try a smaller CSV/.xlsx file, or add supervisors manually.",
                )
                return redirect("upload_supervisors")

            # Return credentials as downloadable CSV.
            out = io.StringIO()
            writer = csv.DictWriter(
                out, fieldnames=["full_name", "username", "email", "password", "school", "note"]
            )
            writer.writeheader()
            for item in created_credentials:
                writer.writerow(item)

            resp = HttpResponse(out.getvalue(), content_type="text/csv")
            resp["Content-Disposition"] = 'attachment; filename="supervisors_login_details.csv"'
            return resp
    else:
        form = UploadFileForm()

    return render(request, "overall/upload_supervisors.html", {"form": form})


@login_required
def overall_add_supervisor(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    created_credentials = None
    if request.method == "POST":
        form = ManualSupervisorForm(request.POST)
        if form.is_valid():
            try:
                created_credentials = _provision_school_supervisor(
                    school=profile.school,
                    full_name=form.cleaned_data["full_name"],
                    email=form.cleaned_data.get("email") or "",
                    username=form.cleaned_data.get("username") or "",
                    password=form.cleaned_data.get("password") or None,
                )
            except ValueError as exc:
                messages.error(request, str(exc))
            except Exception as exc:
                from django.db import IntegrityError

                if isinstance(exc, IntegrityError) and "username" in str(exc).lower():
                    messages.error(
                        request,
                        "That username is already taken. Leave username blank to auto-generate a unique login.",
                    )
                else:
                    raise
            else:
                messages.success(
                    request,
                    f"Supervisor {created_credentials['full_name']} added. "
                    f"Username: {created_credentials['username']}",
                )
                form = ManualSupervisorForm()
    else:
        form = ManualSupervisorForm()

    return render(
        request,
        "overall/add_supervisor.html",
        {"form": form, "created_credentials": created_credentials},
    )


@login_required
def overall_add_student(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    supervisor_id = request.GET.get("supervisor")
    initial = {}
    if supervisor_id and supervisor_id.isdigit():
        if UserProfile.objects.filter(
            school=profile.school,
            role=Role.SUPERVISOR,
            user_id=int(supervisor_id),
        ).exists():
            initial["supervisor"] = int(supervisor_id)

    if request.method == "POST":
        form = OverallStudentForm(request.POST, school=profile.school)
        if form.is_valid():
            student = form.save(commit=False)
            student.school = profile.school
            student.save()
            messages.success(request, f"Student {student.full_name} added successfully.")
            if student.supervisor_id:
                return redirect(
                    reverse(
                        "overall_supervisor_students",
                        kwargs={"supervisor_id": student.supervisor_id},
                    )
                )
            return redirect("overall_students_unassigned")
    else:
        form = OverallStudentForm(initial=initial, school=profile.school)

    has_supervisors = UserProfile.objects.filter(
        school=profile.school, role=Role.SUPERVISOR
    ).exists()

    return render(
        request,
        "overall/add_student.html",
        {
            "form": form,
            "has_supervisors": has_supervisors,
        },
    )


@login_required
def upload_students(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    supervisor_memberships = list(
        UserProfile.objects.select_related("user").filter(school=profile.school, role=Role.SUPERVISOR)
    )
    supervisor_by_name: dict[str, object] = {}
    supervisor_by_username: dict[str, object] = {}
    supervisor_display_names: list[str] = []
    for membership in supervisor_memberships:
        user = membership.user
        display_name = (user.first_name or "").strip() or user.username
        supervisor_display_names.append(display_name)
        name_key = " ".join(display_name.lower().split())
        if name_key and name_key not in supervisor_by_name:
            supervisor_by_name[name_key] = user
        supervisor_by_username[user.username.lower()] = user
    supervisor_display_names = sorted(supervisor_display_names, key=str.casefold)

    if request.method == "POST":
        form = UploadFileForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                rows = _read_uploaded_table(form.cleaned_data["file"], upload_kind="student")
            except ValueError as exc:
                messages.error(request, str(exc))
                return redirect("upload_students")
            except Exception:
                messages.error(
                    request,
                    "Could not read that file. Upload a CSV or Excel (.xlsx) file and try again.",
                )
                return redirect("upload_students")

            if not rows:
                messages.error(request, "File has no rows.")
                return redirect("upload_students")

            if not supervisor_memberships:
                messages.error(
                    request,
                    "No supervisors found. Upload supervisors first, then upload students.",
                )
                return redirect("upload_students")

            created = 0
            try:
                with transaction.atomic():
                    for idx, row in enumerate(rows, start=1):
                        full_name = row_lookup(
                            row,
                            "full_name",
                            "name",
                            "student_name",
                            "student",
                            "learner_name",
                            allow_lone_value=True,
                        )[:200]
                        student_no = row_lookup(row, "student_no", "reg_no", "registration_no", "admission_no")
                        class_level = row_lookup(row, "class_level", "class", "form", "level", "grade").upper()
                        stream = row_lookup(row, "stream", "house")
                        supervisor_ref = row_lookup(
                            row,
                            "supervisor_name",
                            "supervisor_full_name",
                            "supervisor",
                            "supervisor_username",
                            "username",
                        )

                        if not full_name:
                            raise ValueError(
                                f"Row {idx}: missing student name. "
                                f"Add a column named full_name or Name. "
                                f"Columns detected: {headers_found(row)}"
                            )
                        if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
                            raise ValueError(
                                f"Row {idx}: invalid class '{class_level}' (use S1, S2, S3, S4, S5, or S6)"
                            )
                        if not supervisor_ref:
                            raise ValueError(
                                f"Row {idx}: missing supervisor. "
                                f"Add a supervisor_name column with the supervisor's full name "
                                f"(must match a supervisor already uploaded)."
                            )

                        supervisor_user = supervisor_by_name.get(
                            " ".join(supervisor_ref.lower().split())
                        ) or supervisor_by_username.get(supervisor_ref.strip().lower())
                        if not supervisor_user:
                            raise ValueError(
                                f"Row {idx}: supervisor '{supervisor_ref}' not found. "
                                f"Use the supervisor's full name exactly as shown on the supervisors list."
                            )

                        Student.objects.create(
                            school=profile.school,
                            full_name=full_name,
                            student_no=student_no,
                            class_level=class_level,
                            stream=stream,
                            supervisor=supervisor_user,
                        )
                        created += 1
            except ValueError as exc:
                messages.error(request, str(exc))
                return redirect("upload_students")
            except Exception:
                messages.error(
                    request,
                    "Student upload failed unexpectedly. Check the file columns and try a smaller batch.",
                )
                return redirect("upload_students")

            messages.success(request, f"Uploaded {created} students.")
            return redirect("overall_dashboard")
    else:
        form = UploadFileForm()

    return render(
        request,
        "overall/upload_students.html",
        {
            "form": UploadFileForm(),
            "supervisor_names": supervisor_display_names,
        },
    )


@login_required
def supervisor_dashboard(request):
    profile = _require_profile(request)
    if not _is_supervisor(profile):
        raise Http404()

    assigned_classes = _supervisor_student_classes(request.user, profile.school)
    class_filter = _resolve_supervisor_class_filter(request, assigned_classes)

    students_qs = Student.objects.active_only().filter(
        school=profile.school, supervisor=request.user
    ).order_by(Lower("full_name"))
    if class_filter:
        students_qs = students_qs.filter(class_level=class_filter)
    students = list(students_qs)

    year, term = _score_period(request)
    checklist_class_levels = active_checklist_class_levels(profile.school_id, year=year)
    has_checklist = bool(checklist_class_levels)
    theme_classes = [class_filter] if class_filter else assigned_classes
    class_project_themes = project_themes_by_class(
        profile.school_id, year=year, class_levels=theme_classes
    )
    scheme = _active_scheme()
    competencies = list(Competency.objects.filter(scheme=scheme).order_by("order", "id")) if scheme else []
    assigned_ids = list(
        Student.objects.active_only()
        .filter(school=profile.school, supervisor=request.user)
        .values_list("id", flat=True)
    )
    marked_terms = terms_with_marks_for_students(assigned_ids, year)
    if scheme and students:
        score_maps = student_display_score_maps(
            students,
            scheme,
            year,
            term,
            checklist_for=lambda s: resolve_checklist_for_student(s, year=year, term=term),
        )
        student_rows = [
            {"student": s, "score_map": score_maps.get(s.id, {})} for s in students
        ]
    else:
        student_rows = [{"student": s, "score_map": {}} for s in students]

    return render(
        request,
        "supervisor/dashboard.html",
        {
            "students": students,
            "student_rows": student_rows,
            "competencies": competencies,
            "has_observation_checklist": has_checklist,
            "checklist_class_levels": checklist_class_levels,
            "class_project_themes": class_project_themes,
            "current_year": year,
            "current_term": term,
            "marks_on_other_terms": sorted(t for t in marked_terms if t != term),
            "current_term_has_marks": term in marked_terms,
            "assigned_classes": assigned_classes,
            "class_filter": class_filter,
        },
    )


@login_required
def upload_evidence(request, student_id: int):
    profile = _require_profile(request)
    if not _is_supervisor(profile):
        raise Http404()

    student = get_object_or_404(Student, id=student_id, school=profile.school, supervisor=request.user)
    evidence_slots = student_evidence_slots(student)

    if request.method == "POST":
        saved = 0
        with transaction.atomic():
            for slot in evidence_slots:
                uploaded = request.FILES.get(slot["field_name"])
                if not uploaded:
                    continue
                ProjectEvidence.objects.update_or_create(
                    student=student,
                    category=slot["category"],
                    defaults={
                        "uploaded_by": request.user,
                        "file": uploaded,
                        "caption": slot["label"],
                    },
                )
                saved += 1

        if saved == 0:
            messages.error(request, "Please select at least one photo to upload or replace.")
        else:
            messages.success(request, f"Saved {saved} evidence photo(s).")
        return redirect("upload_evidence", student_id=student.id)

    return render(
        request,
        "supervisor/upload_evidence.html",
        {
            "student": student,
            "evidence_slots": evidence_slots,
            "evidence_complete": evidence_complete(student),
            "evidence_count": evidence_completion_count(student),
        },
    )


@login_required
def overall_evidence_index(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = (request.GET.get("class") or "").upper()
    students = Student.objects.active_only().filter(school=profile.school).select_related("supervisor").order_by(
        Lower("full_name"), "class_level"
    )
    if class_level in {c for c, _ in SecondaryClassLevel.choices}:
        students = students.filter(class_level=class_level)

    student_ids = list(students.values_list("id", flat=True))
    evidence_counts = {sid: 0 for sid in student_ids}
    if student_ids:
        from django.db.models import Count

        for row in (
            ProjectEvidence.objects.filter(student_id__in=student_ids)
            .values("student_id")
            .annotate(count=Count("id"))
        ):
            evidence_counts[row["student_id"]] = row["count"]

    required = len(EvidenceCategory)
    rows = [
        {
            "student": s,
            "photo_count": evidence_counts.get(s.id, 0),
            "evidence_complete": evidence_counts.get(s.id, 0) >= required,
        }
        for s in students
    ]

    return render(
        request,
        "overall/evidence_index.html",
        {
            "rows": rows,
            "class_level": class_level,
            "class_levels": [c for c, _ in SecondaryClassLevel.choices],
        },
    )


@login_required
def supervisor_evidence_index(request):
    profile = _require_profile(request)
    if not _is_supervisor(profile):
        raise Http404()

    assigned_classes = _supervisor_student_classes(request.user, profile.school)
    class_filter = _resolve_supervisor_class_filter(request, assigned_classes)

    students_qs = Student.objects.active_only().filter(
        school=profile.school, supervisor=request.user
    ).order_by(Lower("full_name"))
    if class_filter:
        students_qs = students_qs.filter(class_level=class_filter)

    required = len(EvidenceCategory)
    rows = [
        {
            "student": s,
            "photo_count": evidence_completion_count(s),
            "evidence_complete": evidence_complete(s),
        }
        for s in students_qs
    ]

    return render(
        request,
        "supervisor/evidence_index.html",
        {
            "rows": rows,
            "required_photos": required,
            "assigned_classes": assigned_classes,
            "class_filter": class_filter,
        },
    )


@login_required
def overall_student_evidence(request, student_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    student = get_object_or_404(Student, id=student_id, school=profile.school)
    return render(
        request,
        "overall/student_evidence.html",
        {
            "student": student,
            "evidence_slots": student_evidence_slots(student),
            "evidence_complete": evidence_complete(student),
            "required_photos": len(EvidenceCategory),
        },
    )


@login_required
def download_evidence(request, evidence_id: int):
    profile = _require_profile(request)
    # Both overall supervisor and supervisors of that student can download.
    if profile.role not in (Role.OVERALL_SUPERVISOR, Role.SUPERVISOR):
        raise Http404()

    ev = get_object_or_404(ProjectEvidence.objects.select_related("student__school", "student__supervisor"), id=evidence_id)
    if ev.student.school_id != profile.school_id:
        raise Http404()
    if profile.role == Role.SUPERVISOR and ev.student.supervisor_id != request.user.id:
        raise Http404()

    return FileResponse(ev.file.open("rb"), as_attachment=True, filename=ev.file.name.split("/")[-1])


def _score_period(request) -> tuple[int, int]:
    """Year and term from the query string, otherwise the current school period."""
    try:
        year = int(request.GET.get("year") or timezone.now().year)
    except (TypeError, ValueError):
        year = timezone.now().year
    if year < 2000 or year > 2100:
        year = timezone.now().year
    try:
        term = int(request.GET.get("term") or current_school_term())
    except (TypeError, ValueError):
        term = current_school_term()
    if term not in (1, 2, 3):
        term = current_school_term()
    return year, term


def _active_scheme(scheme_id: str | None = None) -> AssessmentScheme | None:
    scheme = AssessmentScheme.objects.filter(active=True).order_by("id").first()
    if scheme_id and str(scheme_id).isdigit():
        scheme = AssessmentScheme.objects.filter(id=int(scheme_id)).first() or scheme
    return scheme


def _existing_score_map(student: Student, scheme: AssessmentScheme, year: int, term: int) -> dict[int, Decimal]:
    assessment = ProjectAssessment.objects.filter(
        student=student, scheme=scheme, year=year, term=term
    ).prefetch_related("scores").first()
    if not assessment:
        return {}
    return {sc.competency_id: sc.score for sc in assessment.scores.all()}


def _save_single_competency_score(
    request,
    student: Student,
    scheme: AssessmentScheme,
    year: int,
    term: int,
    competency_id: int,
) -> tuple[bool, str | None, str]:
    """Save one competency percentage (UNEB phased submission)."""
    competency = Competency.objects.filter(id=competency_id, scheme=scheme).first()
    if not competency:
        return False, "Invalid competency for this scheme.", ""

    raw = (request.POST.get(f"score_{competency.id}") or "").strip()
    if raw == "":
        return False, f"Enter a score for {competency.name} before submitting.", competency.name
    try:
        val = Decimal(raw)
        if val < 0 or val > 100:
            return False, f"Score for {competency.name} must be between 0 and 100.", competency.name
    except InvalidOperation:
        return False, f"Invalid score for {competency.name}.", competency.name

    with transaction.atomic():
        assessment, _ = ProjectAssessment.objects.get_or_create(
            student=student,
            scheme=scheme,
            year=year,
            term=term,
            defaults={"created_by": request.user},
        )
        assessment.created_by = request.user
        assessment.save(update_fields=["created_by", "updated_at"])

        now = timezone.now()
        obj, created = CompetencyScore.objects.update_or_create(
            assessment=assessment,
            competency=competency,
            defaults={"score": val, "submitted_at": now},
        )
        if not created:
            obj.score = val
            obj.submitted_at = now
            obj.save(update_fields=["score", "submitted_at"])

    return True, None, competency.name


def _save_student_scores(
    request,
    student: Student,
    scheme: AssessmentScheme,
    year: int,
    term: int,
) -> tuple[bool, str | None]:
    """Returns (success, error_message)."""
    competencies = list(Competency.objects.filter(scheme=scheme).order_by("order", "id"))
    if not competencies:
        return False, "Selected scheme has no competencies configured."

    parsed_scores: dict[int, Decimal] = {}
    for c in competencies:
        key = f"score_{c.id}"
        raw = (request.POST.get(key) or "").strip()
        if raw == "":
            continue
        try:
            val = Decimal(raw)
            if val < 0 or val > 100:
                return False, f"Score for {c.name} must be between 0 and 100."
            parsed_scores[c.id] = val
        except InvalidOperation:
            return False, f"Invalid score for {c.name}."

    with transaction.atomic():
        assessment, _ = ProjectAssessment.objects.get_or_create(
            student=student,
            scheme=scheme,
            year=year,
            term=term,
            defaults={"created_by": request.user},
        )
        assessment.created_by = request.user
        assessment.save(update_fields=["created_by", "updated_at"])

        existing = {s.competency_id: s for s in CompetencyScore.objects.filter(assessment=assessment)}
        for c in competencies:
            if c.id not in parsed_scores:
                continue
            if c.id in existing:
                obj = existing[c.id]
                obj.score = parsed_scores[c.id]
                obj.save(update_fields=["score"])
            else:
                CompetencyScore.objects.create(assessment=assessment, competency=c, score=parsed_scores[c.id])

    return True, None


def _parse_observation_post_for_competency(
    request, checklist: ObservationChecklist, competency_id: int
) -> dict[int, int]:
    from .models import ObservationCriterion

    criteria = ObservationCriterion.objects.filter(
        checklist=checklist, competency_id=competency_id
    )
    checkbox = uses_checkbox_scoring(checklist)
    ratings: dict[int, int] = {}
    if checkbox:
        for c in criteria:
            if c.max_rating > 0:
                ratings[c.id] = 1 if request.POST.get(f"obs_{c.id}") == "1" else 0
    else:
        valid_set = set(criteria.values_list("id", flat=True))
        for key, val in request.POST.items():
            if not key.startswith("obs_"):
                continue
            try:
                cid = int(key[4:])
            except ValueError:
                continue
            if cid not in valid_set:
                continue
            try:
                ratings[cid] = int(val)
            except ValueError:
                continue
    return ratings


def _parse_observation_post(request, checklist: ObservationChecklist) -> dict[int, int]:
    from .models import ObservationCriterion

    valid_set = set(
        ObservationCriterion.objects.filter(checklist=checklist).values_list("id", flat=True)
    )
    checkbox = uses_checkbox_scoring(checklist)
    ratings: dict[int, int] = {}
    if checkbox:
        for cid in valid_set:
            ratings[cid] = 1 if request.POST.get(f"obs_{cid}") == "1" else 0
    else:
        for key, val in request.POST.items():
            if not key.startswith("obs_"):
                continue
            try:
                cid = int(key[4:])
            except ValueError:
                continue
            if cid not in valid_set:
                continue
            try:
                ratings[cid] = int(val)
            except ValueError:
                continue
    return ratings


def _render_score_form(
    request,
    student: Student,
    *,
    template: str,
    page_title: str,
    back_url: str,
    allow_project_title_edit: bool = False,
):
    if request.method == "POST":
        meta_form = ScoreStudentForm(request.POST)
        if meta_form.is_valid():
            scheme = meta_form.cleaned_data["scheme"]
            year = meta_form.cleaned_data["year"]
            term = int(meta_form.cleaned_data["term"])
            if allow_project_title_edit:
                student.project_title = (request.POST.get("project_title") or "").strip()
                student.save(update_fields=["project_title"])

            checklist = resolve_checklist_for_student(student, year=year, term=term)
            use_observation = request.POST.get("scoring_mode") == "observation" and checklist

            if use_observation:
                submit_competency = request.POST.get("submit_competency")
                if not submit_competency or not str(submit_competency).isdigit():
                        messages.error(
                        request,
                        "UNEB phased submission: complete one competency and use its Submit button.",
                    )
                else:
                    competency_id = int(submit_competency)
                    ratings = _parse_observation_post_for_competency(
                        request, checklist, competency_id
                    )
                    if not ratings:
                        messages.error(request, "No observation items found for this competency.")
                    else:
                        try:
                            pct, comp_name = save_competency_observation_ratings_and_score(
                                student=student,
                                checklist=checklist,
                                year=year,
                                term=term,
                                competency_id=competency_id,
                                ratings_by_criterion=ratings,
                                user=request.user,
                                scheme=scheme,
                            )
                        except ValueError as exc:
                            messages.error(request, str(exc))
                        else:
                            if pct is not None:
                                messages.success(
                                    request,
                                    f"{comp_name} submitted successfully — competency score: {pct}%.",
                                )
                            else:
                                messages.warning(
                                    request,
                                    f"{comp_name}: tick at least one observation before submitting.",
                                )
                            q = f"?year={year}&term={term}"
                            if scheme:
                                q += f"&scheme_id={scheme.id}"
                            return redirect(f"{request.path}{q}")
            else:
                submit_competency = request.POST.get("submit_competency")
                if not submit_competency or not str(submit_competency).isdigit():
                    messages.error(
                        request,
                        "UNEB requires phased submission: use Submit on one competency at a time.",
                    )
                else:
                    ok, err, comp_name = _save_single_competency_score(
                        request,
                        student,
                        scheme,
                        year,
                        term,
                        int(submit_competency),
                    )
                    if allow_project_title_edit:
                        student.project_title = (request.POST.get("project_title") or "").strip()
                        student.save(update_fields=["project_title"])
                    if ok:
                        raw = (request.POST.get(f"score_{submit_competency}") or "").strip()
                        messages.success(
                            request,
                            f"{comp_name} submitted successfully — score: {raw}%.",
                        )
                        q = f"?year={year}&term={term}"
                        if scheme:
                            q += f"&scheme_id={scheme.id}"
                        return redirect(f"{request.path}{q}")
                    messages.error(request, err or "Could not save score.")
        else:
            messages.error(
                request,
                "Marks were not saved. Check the year, term, and assessment scheme, then submit again.",
            )
        scheme = (
            meta_form.cleaned_data["scheme"]
            if meta_form.is_valid()
            else _active_scheme(request.POST.get("scheme"))
        )
        year = meta_form.cleaned_data.get("year") if meta_form.is_valid() else int(request.POST.get("year") or timezone.now().year)
        term = (
            int(meta_form.cleaned_data["term"])
            if meta_form.is_valid()
            else int(request.POST.get("term") or current_school_term())
        )
    else:
        year = int(request.GET.get("year") or timezone.now().year)
        term = int(request.GET.get("term") or current_school_term())
        scheme = _active_scheme(request.GET.get("scheme_id"))
        meta_form = ScoreStudentForm(initial={"scheme": scheme, "year": year, "term": str(term)})

    competencies = list(Competency.objects.filter(scheme=scheme).order_by("order", "id")) if scheme else {}

    checklist = None
    observation_structure = []
    observation_ratings = {}
    computed_percentages = {}
    submitted_competency_ids: set[int] = set()
    checkbox_mode = True
    assessment = None
    if scheme:
        assessment = ProjectAssessment.objects.filter(
            student=student, scheme=scheme, year=year, term=term
        ).prefetch_related("scores").first()
        checklist = resolve_checklist_for_student(student, year=year, term=term)
        if checklist:
            observation_structure = get_observation_structure(checklist)
            observation_ratings = existing_ratings_map(student.id, checklist.id, year, term)
            checkbox_mode = uses_checkbox_scoring(checklist)
            submitted_competency_ids = get_submitted_competency_ids(
                student.id, checklist, year, term
            )
            if observation_ratings:
                existing_comp_ids = submitted_competency_ids
                computed_percentages = percentages_to_persist(
                    checklist, observation_ratings, existing_comp_ids
                )
            else:
                computed_percentages = {}
        elif assessment:
            submitted_competency_ids = submitted_competency_ids_for_assessment(assessment)

    existing_scores = (
        student_display_score_maps(
            [student],
            scheme,
            year,
            term,
            checklist_for=lambda _student: checklist,
        ).get(student.id, {})
        if scheme
        else {}
    )
    # A visible saved percentage is a submission, even if it was stored under another scheme.
    submitted_competency_ids = set(submitted_competency_ids) | set(existing_scores)

    obs_template = "shared/score_student_observation.html"
    use_template = obs_template if checklist and observation_structure else template

    active_profile = get_active_profile(request)
    observation_pdf_url = ""
    if (
        checklist
        and observation_structure
        and observation_ratings
        and active_profile
        and _is_overall(active_profile)
    ):
        observation_pdf_url = (
            f"{reverse('export_student_observation_checklist_pdf', kwargs={'student_id': student.id})}"
            f"?year={year}&term={term}&scheme_id={scheme.id if scheme else ''}"
        )

    return render(
        request,
        use_template,
        {
            "student": student,
            "meta_form": meta_form,
            "scheme": scheme,
            "competencies": competencies,
            "existing_scores": existing_scores,
            "page_title": page_title,
            "back_url": back_url,
            "allow_project_title_edit": allow_project_title_edit,
            "checklist": checklist,
            "observation_structure": observation_structure,
            "observation_ratings": observation_ratings,
            "computed_percentages": computed_percentages,
            "rating_labels": OBSERVATION_RATING_LABELS_CHECKBOX
            if checkbox_mode
            else OBSERVATION_RATING_LABELS_SCALE,
            "checkbox_mode": checkbox_mode,
            "use_observation_scoring": bool(checklist and observation_structure),
            "checklist_document": _checklist_document_context(checklist),
            "checklist_competency_headings": _checklist_competency_headings(checklist),
            "class_project_theme": get_student_project_theme(
                student,
                year=year,
                term=term,
            ),
            "observation_pdf_url": observation_pdf_url,
            # Overall supervisors and assigned supervisors can both edit scores.
            "read_only_observation": False,
            "year": year,
            "term": term,
            "submitted_competency_ids": submitted_competency_ids,
        },
    )


def _checklist_document_context(checklist=None):
    return checklist_document_for(checklist)


def _checklist_competency_headings(checklist=None):
    return checklist_headings_for(checklist)


@login_required
def score_student(request, student_id: int):
    profile = _require_profile(request)
    if not _is_supervisor(profile):
        raise Http404()

    student = get_object_or_404(Student, id=student_id, school=profile.school, supervisor=request.user)
    back_url = reverse("supervisor_dashboard")
    if student.class_level:
        back_url = f"{back_url}?class={student.class_level}"
    return _render_score_form(
        request,
        student,
        template="shared/score_student_form.html",
        page_title="Score student",
        back_url=back_url,
        allow_project_title_edit=True,
    )


@login_required
def overall_edit_student_scores(request, student_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    student = get_object_or_404(Student, id=student_id, school=profile.school)
    year = request.GET.get("year") or request.POST.get("year") or timezone.now().year
    term = request.GET.get("term") or request.POST.get("term") or current_school_term()
    scheme_id = request.GET.get("scheme_id") or request.POST.get("scheme")
    back_url = (
        f"{reverse('class_scores', kwargs={'class_level': student.class_level})}"
        f"?year={year}&term={term}&scheme_id={scheme_id or ''}"
    )
    return _render_score_form(
        request,
        student,
        template="shared/score_student_form.html",
        page_title="Edit student scores",
        back_url=back_url,
    )


@login_required
def overall_edit_student(request, student_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    student = get_object_or_404(Student, id=student_id, school=profile.school)
    year = request.GET.get("year") or request.POST.get("year") or timezone.now().year
    term = request.GET.get("term") or request.POST.get("term") or current_school_term()
    scheme_id = request.GET.get("scheme_id") or request.POST.get("scheme_id") or ""
    back_url = (
        f"{reverse('class_scores', kwargs={'class_level': student.class_level})}"
        f"?year={year}&term={term}&scheme_id={scheme_id}"
    )

    if request.method == "POST":
        form = OverallStudentForm(request.POST, instance=student, school=profile.school)
        if form.is_valid():
            updated = form.save()
            messages.success(request, f"Updated details for {updated.full_name}.")
            new_class = updated.class_level
            return redirect(
                f"{reverse('class_scores', kwargs={'class_level': new_class})}"
                f"?year={year}&term={term}&scheme_id={scheme_id}"
            )
    else:
        form = OverallStudentForm(instance=student, school=profile.school)

    return render(
        request,
        "overall/edit_student.html",
        {
            "student": student,
            "form": form,
            "back_url": back_url,
            "year": year,
            "term": term,
            "scheme_id": scheme_id,
        },
    )


@login_required
def class_scores(request, class_level: str):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()

    form = ClassFilterForm(
        request.GET or None,
        initial={"class_level": class_level, "year": timezone.now().year, "term": str(current_school_term())},
    )
    if not form.is_valid():
        form = ClassFilterForm(initial={"class_level": class_level, "year": timezone.now().year, "term": str(current_school_term())})

    year, term = _score_period(request)
    scheme = AssessmentScheme.objects.filter(active=True).order_by("id").first()
    scheme_id = request.GET.get("scheme_id")
    if scheme_id and scheme_id.isdigit():
        scheme = AssessmentScheme.objects.filter(id=int(scheme_id)).first() or scheme

    search_q = (request.GET.get("q") or "").strip()
    students_qs = Student.objects.active_only().filter(
        school=profile.school, class_level=class_level
    ).select_related("supervisor").order_by(Lower("full_name"))
    total_students = students_qs.count()
    if search_q:
        students_qs = students_qs.filter(
            Q(full_name__icontains=search_q)
            | Q(student_no__icontains=search_q)
            | Q(project_title__icontains=search_q)
            | Q(stream__icontains=search_q)
            | Q(supervisor__username__icontains=search_q)
            | Q(supervisor__first_name__icontains=search_q)
            | Q(supervisor__last_name__icontains=search_q)
        )
    students = list(students_qs)
    class_student_ids = list(
        Student.objects.active_only()
        .filter(school=profile.school, class_level=class_level)
        .values_list("id", flat=True)
    )
    marked_terms = terms_with_marks_for_students(class_student_ids, year)
    assessments = (
        ProjectAssessment.objects.select_related("student")
        .filter(student__in=students, scheme=scheme, year=year, term=term)
        .prefetch_related("scores__competency")
    )
    by_student = {a.student_id: a for a in assessments}

    competencies = list(Competency.objects.filter(scheme=scheme).order_by("order", "id")) if scheme else []
    checklist = (
        resolve_checklist(profile.school_id, year=year, term=term, class_level=class_level)
        if scheme
        else None
    )
    score_maps = (
        student_display_score_maps(
            students,
            scheme,
            year,
            term,
            checklist_for=lambda _student: checklist,
        )
        if scheme
        else {}
    )

    table = []
    for s in students:
        a = by_student.get(s.id)
        has_obs_pdf = (
            bool(checklist)
            and student_has_observation_ratings(s.id, checklist.id, year, term)
        )
        table.append(
            {
                "student": s,
                "assessment": a,
                "score_map": score_maps.get(s.id, {}),
                "has_observation_pdf": has_obs_pdf,
            }
        )

    any_observation_pdf = any(row["has_observation_pdf"] for row in table)

    return render(
        request,
        "overall/class_scores.html",
        {
            "class_level": class_level,
            "scheme": scheme,
            "year": year,
            "term": term,
            "competencies": competencies,
            "table": table,
            "schemes": AssessmentScheme.objects.all().order_by("name"),
            "checklist": checklist,
            "any_observation_pdf": any_observation_pdf,
            "search_q": search_q,
            "total_students": total_students,
            "filtered_students": len(table),
            "marks_on_other_terms": sorted(t for t in marked_terms if t != term),
            "current_term_has_marks": term in marked_terms,
        },
    )


def _build_student_scored_observation_pdf_export(
    *,
    school,
    student: Student,
    year: int,
    term: int,
) -> tuple[bytes, str] | None:
    """Return (pdf_bytes, filename) for a scored checklist, or None if not available."""
    checklist = resolve_checklist_for_student(student, year=year, term=term)
    if not checklist:
        return None

    observation_ratings = existing_ratings_map(student.id, checklist.id, year, term)
    if not observation_ratings:
        return None

    structure = get_observation_structure(checklist)
    assessment = ProjectAssessment.objects.filter(
        student=student, scheme=checklist.scheme, year=year, term=term
    ).prefetch_related("scores").first()
    existing_comp_ids = (
        {s.competency_id for s in assessment.scores.all()} if assessment else set()
    )
    percentages = percentages_to_persist(checklist, observation_ratings, existing_comp_ids)

    theme_obj = get_student_project_theme(student, year=year, term=term)
    theme_text = theme_obj.theme if theme_obj else ""

    pdf_bytes = build_student_scored_observation_pdf(
        school=school,
        student=student,
        checklist=checklist,
        year=year,
        term=term,
        observation_structure=structure,
        observation_ratings=observation_ratings,
        competency_percentages=percentages,
        competency_headings=_checklist_competency_headings(checklist),
        checklist_document=_checklist_document_context(checklist),
        project_theme=theme_text,
        assessed_by=latest_assessor_display_name(student.id, checklist.id, year, term),
    )

    safe_name = _safe_folder_name(student.full_name)[:40]
    student_no = re.sub(r"[^\w-]", "", student.student_no or "")[:20]
    suffix = f"_{student_no}" if student_no else ""
    filename = (
        f"{school.name}_{safe_name}{suffix}_observation_scored_{year}_T{term}.pdf"
    ).replace(" ", "_")
    return pdf_bytes, filename


@login_required
def export_student_observation_checklist_pdf(request, student_id: int):
    """Download the learner's scored observation checklist (not the blank UNEB upload)."""
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    student = get_object_or_404(
        Student.objects.select_related("supervisor", "school"),
        id=student_id,
        school=profile.school,
    )
    year = int(request.GET.get("year") or timezone.now().year)
    term = int(request.GET.get("term") or current_school_term())

    result = _build_student_scored_observation_pdf_export(
        school=profile.school,
        student=student,
        year=year,
        term=term,
    )
    if not result:
        checklist = resolve_checklist_for_student(student, year=year, term=term)
        if not checklist:
            messages.error(request, "No observation checklist for this class and year.")
        else:
            messages.error(
                request,
                "This learner has not been scored on the observation checklist yet.",
            )
        q = f"?year={year}&term={term}"
        scheme_id = request.GET.get("scheme_id")
        if scheme_id:
            q += f"&scheme_id={scheme_id}"
        return redirect(
            f"{reverse('class_scores', kwargs={'class_level': student.class_level})}{q}"
        )

    pdf_bytes, filename = result
    resp = HttpResponse(pdf_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


@login_required
def export_class_observation_checklists_zip(request, class_level: str):
    """ZIP of scored observation checklist PDFs for all learners in the class who have been assessed."""
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()

    year = int(request.GET.get("year") or timezone.now().year)
    term = int(request.GET.get("term") or current_school_term())
    scheme_id = request.GET.get("scheme_id")
    q = f"?year={year}&term={term}"
    if scheme_id:
        q += f"&scheme_id={scheme_id}"
    back_url = f"{reverse('class_scores', kwargs={'class_level': class_level})}{q}"

    checklist = resolve_checklist(
        profile.school_id, year=year, term=term, class_level=class_level
    )
    if not checklist:
        messages.error(request, "No observation checklist for this class and year.")
        return redirect(back_url)

    students = list(
        Student.objects.active_only().filter(school=profile.school, class_level=class_level)
        .select_related("supervisor", "school")
        .order_by(Lower("full_name"))
    )

    zip_buf = io.BytesIO()
    included = 0
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for student in students:
            result = _build_student_scored_observation_pdf_export(
                school=profile.school,
                student=student,
                year=year,
                term=term,
            )
            if not result:
                continue
            pdf_bytes, filename = result
            zf.writestr(f"scored_checklists/{filename}", pdf_bytes)
            included += 1

    if included == 0:
        messages.error(
            request,
            "No learners in this class have been scored on the observation checklist yet.",
        )
        return redirect(back_url)

    zip_buf.seek(0)
    base = (
        f"{profile.school.name}_{class_level}_{year}_T{term}_scored_observation_checklists"
    ).replace(" ", "_")
    resp = HttpResponse(zip_buf.getvalue(), content_type="application/zip")
    resp["Content-Disposition"] = f'attachment; filename="{base}.zip"'
    return resp


@login_required
def export_class_scores(request, class_level: str):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()

    year = int((request.GET.get("year") or timezone.now().year))
    term = int((request.GET.get("term") or current_school_term()))
    scheme = AssessmentScheme.objects.filter(active=True).order_by("id").first()
    scheme_id = request.GET.get("scheme_id")
    if scheme_id and scheme_id.isdigit():
        scheme = AssessmentScheme.objects.filter(id=int(scheme_id)).first() or scheme
    if not scheme:
        raise Http404("No scheme configured")

    competencies = list(Competency.objects.filter(scheme=scheme).order_by("order", "id"))
    students = list(
        Student.objects.active_only()
        .filter(school=profile.school, class_level=class_level)
        .order_by(Lower("full_name"))
    )

    wb = Workbook()
    ws = wb.active
    ws.title = f"{class_level} Scores"

    headers = ["Student Name", "Supervisor", "Student No", "Stream"] + [c.name for c in competencies]
    ws.append(headers)

    checklist = resolve_checklist(profile.school_id, year=year, term=term, class_level=class_level)
    score_maps = student_display_score_maps(
        students,
        scheme,
        year,
        term,
        checklist_for=lambda _student: checklist,
    )
    for s in students:
        score_map = score_maps.get(s.id, {})
        row = [
            s.full_name,
            s.supervisor.get_full_name() or s.supervisor.username,
            s.student_no,
            s.stream,
        ]
        row += [float(score_map[c.id]) if c.id in score_map else "" for c in competencies]
        ws.append(row)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    filename = f"{profile.school.name}_{class_level}_{scheme.name}_{year}_T{term}.xlsx".replace(" ", "_")
    resp = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


@login_required
def export_class_competency_scores(request, class_level: str):
    """Export scores for a single competency (UNEB phased submission file)."""
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()

    competency_id = request.GET.get("competency_id")
    if not competency_id or not str(competency_id).isdigit():
        raise Http404("Select a competency to export.")

    year = int((request.GET.get("year") or timezone.now().year))
    term = int((request.GET.get("term") or current_school_term()))
    scheme = AssessmentScheme.objects.filter(active=True).order_by("id").first()
    scheme_id = request.GET.get("scheme_id")
    if scheme_id and scheme_id.isdigit():
        scheme = AssessmentScheme.objects.filter(id=int(scheme_id)).first() or scheme
    if not scheme:
        raise Http404("No scheme configured")

    competency = get_object_or_404(Competency, pk=int(competency_id), scheme=scheme)

    students = list(
        Student.objects.active_only()
        .filter(school=profile.school, class_level=class_level)
        .select_related("supervisor")
        .order_by(Lower("full_name"))
    )
    assessments = (
        ProjectAssessment.objects.filter(student__in=students, scheme=scheme, year=year, term=term)
        .prefetch_related("scores")
    )
    by_student = {a.student_id: a for a in assessments}
    checklist = resolve_checklist(profile.school_id, year=year, term=term, class_level=class_level)
    score_maps = student_display_score_maps(
        students,
        scheme,
        year,
        term,
        checklist_for=lambda _student: checklist,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = competency.name[:31]
    ws.append(
        [
            "Student Name",
            "Supervisor",
            "Student No",
            "Stream",
            f"{competency.name} (%)",
            "Submitted",
        ]
    )

    for s in students:
        a = by_student.get(s.id)
        score_map = score_maps.get(s.id, {})
        submitted = ""
        if a:
            cs = CompetencyScore.objects.filter(assessment=a, competency=competency).first()
            if cs and cs.submitted_at:
                submitted = cs.submitted_at.strftime("%Y-%m-%d %H:%M")
        row_score = score_map.get(competency.id)
        ws.append(
            [
                s.full_name,
                s.supervisor.get_full_name() or s.supervisor.username if s.supervisor_id else "",
                s.student_no,
                s.stream,
                float(row_score) if row_score is not None and row_score != "" else "",
                submitted,
            ]
        )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    safe_comp = re.sub(r"[^\w]+", "_", competency.name).strip("_")[:40]
    filename = (
        f"{profile.school.name}_{class_level}_{safe_comp}_{year}_T{term}.xlsx".replace(" ", "_")
    )
    resp = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _class_scores_queryset(profile, class_level: str, year: int, term: int, scheme):
    students = list(
        Student.objects.active_only().filter(school=profile.school, class_level=class_level)
        .select_related("supervisor")
        .order_by(Lower("full_name"))
    )
    competencies = list(Competency.objects.filter(scheme=scheme).order_by("order", "id"))
    assessments = ProjectAssessment.objects.filter(
        student__in=students, scheme=scheme, year=year, term=term
    ).prefetch_related("scores")
    by_student = {a.student_id: a for a in assessments}
    return students, competencies, by_student


def _resolve_class_export(request, profile, class_level: str):
    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()
    year = int(request.GET.get("year") or timezone.now().year)
    term = int(request.GET.get("term") or current_school_term())
    scheme = AssessmentScheme.objects.filter(active=True).order_by("id").first()
    scheme_id = request.GET.get("scheme_id")
    if scheme_id and str(scheme_id).isdigit():
        scheme = AssessmentScheme.objects.filter(id=int(scheme_id)).first() or scheme
    if not scheme:
        raise Http404("No scheme configured")
    return class_level, year, term, scheme


@login_required
def overall_class_project_themes(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    school = profile.school
    year = int(request.GET.get("year") or request.POST.get("year") or timezone.now().year)

    if request.method == "POST":
        action = request.POST.get("action")
        if action == "delete_cohort":
            theme_id = request.POST.get("theme_id")
            if theme_id and str(theme_id).isdigit():
                deleted = CohortProjectTheme.objects.filter(
                    pk=int(theme_id), school=school
                ).delete()[0]
                if deleted:
                    messages.success(request, "UNEB cohort theme removed.")
                else:
                    messages.error(request, "Theme not found.")
            return redirect(f"{reverse('overall_class_project_themes')}?year={year}")

        if action == "delete_class":
            theme_id = request.POST.get("theme_id")
            if theme_id and str(theme_id).isdigit():
                deleted = ClassProjectTheme.objects.filter(
                    pk=int(theme_id), school=school
                ).delete()[0]
                if deleted:
                    messages.success(request, "Class theme removed.")
                else:
                    messages.error(request, "Theme not found.")
            return redirect(f"{reverse('overall_class_project_themes')}?year={year}")

        form_kind = request.POST.get("form_kind", "cohort")
        if form_kind == "class":
            form = ClassProjectThemeForm(request.POST)
            cohort_form = CohortProjectThemeForm(initial={"cohort_year": year})
            if form.is_valid():
                data = form.cleaned_data
                ClassProjectTheme.objects.update_or_create(
                    school=school,
                    class_level=data["class_level"],
                    year=data["year"],
                    defaults={
                        "theme": data["theme"],
                        "updated_by": request.user,
                    },
                )
                messages.success(
                    request,
                    f'Project theme saved for {data["class_level"]} ({data["year"]}).',
                )
                return redirect(f"{reverse('overall_class_project_themes')}?year={year}")
            messages.error(request, "Could not save class theme. Please fix the errors below.")
        else:
            cohort_form = CohortProjectThemeForm(request.POST)
            form = ClassProjectThemeForm(initial={"year": year})
            if cohort_form.is_valid():
                data = cohort_form.cleaned_data
                CohortProjectTheme.objects.update_or_create(
                    school=school,
                    cohort_year=data["cohort_year"],
                    defaults={
                        "theme": data["theme"],
                        "updated_by": request.user,
                    },
                )
                messages.success(
                    request,
                    f'UNEB theme saved for cohort {data["cohort_year"]} '
                    f'(S.3 Term 1 {data["cohort_year"]} → S.4 Term 3 {data["cohort_year"] + 1}).',
                )
                return redirect(f"{reverse('overall_class_project_themes')}?year={year}")
            messages.error(request, "Could not save UNEB theme. Please fix the errors below.")
    else:
        cohort_form = CohortProjectThemeForm(initial={"cohort_year": year})
        form = ClassProjectThemeForm(initial={"year": year})

    active_cohort_years = {year, year - 1}
    cohort_themes = list(
        CohortProjectTheme.objects.filter(
            school=school, cohort_year__in=active_cohort_years
        )
        .select_related("updated_by")
        .order_by("-cohort_year")
    )
    class_themes = list(
        ClassProjectTheme.objects.filter(school=school, year=year)
        .exclude(class_level__in=[SecondaryClassLevel.S3, SecondaryClassLevel.S4])
        .select_related("updated_by")
        .order_by("class_level")
    )

    return render(
        request,
        "overall/class_project_themes.html",
        {
            "cohort_form": cohort_form,
            "form": form,
            "cohort_themes": cohort_themes,
            "class_themes": class_themes,
            "year": year,
        },
    )


@login_required
def overall_upload_logo(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    school = profile.school
    if request.method == "POST":
        form = SchoolBrandingForm(request.POST, request.FILES, instance=school)
        if form.is_valid():
            form.save()
            messages.success(
                request,
                "School profile updated (logo, theme, motto, vision, mission, and home page).",
            )
            return redirect("home")
    else:
        form = SchoolBrandingForm(instance=school)

    return render(
        request,
        "overall/school_branding.html",
        {"school": school, "form": form, "theme_previews": THEMES},
    )


@login_required
def overall_observation_checklists(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    school = profile.school
    from django.db.models import Count

    checklists = list(
        ObservationChecklist.objects.filter(school=school)
        .select_related("scheme", "uploaded_by")
        .annotate(criteria_count=Count("criteria"))
        .order_by("-year", "-created_at")
    )

    if request.method == "POST":
        form = ObservationChecklistUploadForm(request.POST, request.FILES)
        if form.is_valid():
            scheme = _active_scheme()
            if not scheme:
                messages.error(request, "No active assessment scheme. Run set_uganda_competencies first.")
            else:
                cl: ObservationChecklist = form.save(commit=False)
                cl.school = school
                cl.scheme = scheme
                cl.uploaded_by = request.user
                cl.active = True
                if form.cleaned_data.get("use_builtin_checklist"):
                    cl.format_year = form.cleaned_data.get("format_year")
                uploaded_pdf = request.FILES.get("pdf_file")
                if not uploaded_pdf and form.cleaned_data.get("use_builtin_checklist"):
                    builtin = get_builtin_format(cl.format_year)
                    if builtin and builtin.pdf_path and builtin.pdf_path.exists():
                        with builtin.pdf_path.open("rb") as fh:
                            cl.pdf_file.save(builtin.pdf_filename, File(fh), save=False)
                cl.save()
                deactivate_qs = ObservationChecklist.objects.filter(
                    school=school, class_level=cl.class_level, active=True
                ).exclude(pk=cl.pk)
                if cl.cohort_year:
                    deactivate_qs = deactivate_qs.filter(cohort_year=cl.cohort_year)
                else:
                    deactivate_qs = deactivate_qs.filter(
                        year=cl.year, cohort_year__isnull=True
                    )
                deactivate_qs.update(active=False)
                structure_file = request.FILES.get("structure_file")
                if structure_file:
                    n = import_structure_file(cl, structure_file)
                    messages.success(
                        request,
                        f'Checklist "{cl.title}" uploaded with {n} observations from your structure file.',
                    )
                elif form.cleaned_data.get("use_builtin_checklist"):
                    builtin = get_builtin_format(cl.format_year)
                    n = seed_checklist_criteria(cl)
                    label = builtin.label if builtin else "built-in"
                    messages.success(
                        request,
                        f'Checklist "{cl.title}" uploaded with {n} observations ({label}).',
                    )
                else:
                    messages.warning(
                        request,
                        "PDF saved. Upload an Excel/CSV structure file or enable the built-in checklist.",
                    )
                return redirect("overall_observation_checklists")
        messages.error(request, "Could not upload checklist. Please fix the errors below.")
    else:
        form = ObservationChecklistUploadForm(
            initial={
                "year": timezone.now().year,
                "cohort_year": timezone.now().year,
                "builtin_checklist": str(timezone.now().year)
                if get_builtin_for_cohort(timezone.now().year)
                else "",
                "title": get_builtin_for_cohort(timezone.now().year).default_title
                if get_builtin_for_cohort(timezone.now().year)
                else "",
            }
        )

    calendar_year = timezone.now().year
    builtin_meta = {
        str(year): {
            "label": b.label,
            "title": b.default_title,
            "has_pdf": bool(b.pdf_path and b.pdf_path.exists()),
        }
        for year, b in BUILTIN_CHECKLISTS.items()
    }

    return render(
        request,
        "overall/observation_checklists.html",
        {
            "form": form,
            "checklists": checklists,
            "builtin_checklists": list(BUILTIN_CHECKLISTS.values()),
            "cohort_upload_guide": cohort_upload_guide(calendar_year=calendar_year),
            "calendar_year": calendar_year,
            "builtin_meta_json": json.dumps(builtin_meta),
            "builtin_labels": {year: b.label for year, b in BUILTIN_CHECKLISTS.items()},
            "builtin_format_years": builtin_format_years(),
        },
    )


@login_required
def delete_observation_checklist(request, checklist_id: int):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()
    if request.method != "POST":
        return redirect("overall_observation_checklists")

    school = profile.school
    cl = get_object_or_404(ObservationChecklist, pk=checklist_id, school=school)
    label = str(cl)
    if cl.pdf_file:
        cl.pdf_file.delete(save=False)
    cl.delete()
    messages.success(request, f'Observation checklist "{label}" has been deleted.')
    return redirect("overall_observation_checklists")


@login_required
def download_observation_structure_template(request):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    from .builtin_checklists import get_builtin_format, structured_checklist_for
    from .observation_import import builtin_structure_xlsx_filename, export_structure_xlsx_bytes

    format_year_raw = request.GET.get("format") or request.GET.get("cohort")
    if format_year_raw and str(format_year_raw).isdigit():
        format_year = int(format_year_raw)
    else:
        format_year = 2026

    builtin = get_builtin_format(format_year)
    structure = structured_checklist_for(format_year=format_year)
    resp = HttpResponse(
        export_structure_xlsx_bytes(structure),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    filename = (
        builtin_structure_xlsx_filename(format_year)
        if builtin
        else f"observation_checklist_structure_{format_year}.xlsx"
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


@login_required
def download_observation_checklist(request, checklist_id: int):
    profile = _require_profile(request)
    school = profile.school
    if not school:
        raise Http404()

    cl = get_object_or_404(ObservationChecklist, pk=checklist_id, school=school)
    if not cl.pdf_file:
        raise Http404("No PDF file")
    filename = cl.pdf_file.name.split("/")[-1]
    return FileResponse(cl.pdf_file.open("rb"), as_attachment=True, filename=filename)


@login_required
def supervisor_observation_checklists(request):
    profile = _require_profile(request)
    if profile.role not in (Role.SUPERVISOR, Role.OVERALL_SUPERVISOR):
        raise Http404()

    school = profile.school
    year = int(request.GET.get("year") or timezone.now().year)
    supervisor_classes = list(
        Student.objects.active_only()
        .filter(school=school, supervisor=request.user)
        .values_list("class_level", flat=True)
        .distinct()
    )
    cohort_years = {
        cy
        for cl in supervisor_classes
        if (cy := cohort_year_for_class(cl, year)) is not None
    }
    checklists = list(
        ObservationChecklist.objects.filter(school=school, active=True)
        .filter(
            Q(cohort_year__in=cohort_years)
            | Q(year=year, cohort_year__isnull=True)
        )
        .filter(class_level__in=supervisor_classes)
        .exclude(class_level="")
        .order_by("-cohort_year", "class_level", "-created_at")
    )
    class_project_themes = project_themes_by_class(
        school.id, year=year, class_levels=supervisor_classes
    )

    return render(
        request,
        "supervisor/observation_checklists.html",
        {
            "checklists": checklists,
            "year": year,
            "class_project_themes": class_project_themes,
        },
    )


@login_required
def export_class_scores_pdf(request, class_level: str):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level, year, term, scheme = _resolve_class_export(request, profile, class_level)
    students, competencies, by_student = _class_scores_queryset(profile, class_level, year, term, scheme)

    checklist = resolve_checklist(profile.school_id, year=year, term=term, class_level=class_level)
    pdf_bytes = build_class_scores_pdf(
        school=profile.school,
        class_level=class_level,
        year=year,
        term=term,
        scheme_name=scheme.name,
        competencies=competencies,
        students=students,
        by_student=by_student,
        checklist=checklist,
    )
    filename = f"{profile.school.name}_{class_level}_scores_{year}_T{term}.pdf".replace(" ", "_")
    resp = HttpResponse(pdf_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


@login_required
def export_class_project_titles_pdf(request, class_level: str):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level, year, term, _scheme = _resolve_class_export(request, profile, class_level)
    students = list(
        Student.objects.active_only().filter(school=profile.school, class_level=class_level)
        .select_related("supervisor")
        .order_by(Lower("full_name"))
    )

    pdf_bytes = build_class_project_titles_pdf(
        school=profile.school,
        class_level=class_level,
        year=year,
        term=term,
        students=students,
    )
    filename = f"{profile.school.name}_{class_level}_project_titles_{year}_T{term}.pdf".replace(" ", "_")
    resp = HttpResponse(pdf_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


@login_required
def export_class_project_titles(request, class_level: str):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()

    year = int(request.GET.get("year") or timezone.now().year)
    term = int(request.GET.get("term") or current_school_term())

    students = list(
        Student.objects.active_only().filter(school=profile.school, class_level=class_level)
        .select_related("supervisor")
        .order_by(Lower("full_name"))
    )

    wb = Workbook()
    ws = wb.active
    ws.title = f"{class_level} Project Titles"
    ws.append(["Student Name", "Project Title", "Supervisor", "Student No", "Stream"])
    for s in students:
        ws.append(
            [
                s.full_name,
                s.project_title,
                s.supervisor.get_full_name() or s.supervisor.username,
                s.student_no,
                s.stream,
            ]
        )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"{profile.school.name}_{class_level}_project_titles_{year}_T{term}.xlsx".replace(" ", "_")
    resp = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


@login_required
def export_class_evidence_pdf(request, class_level: str):
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()

    students = list(
        Student.objects.active_only().filter(school=profile.school, class_level=class_level)
        .select_related("supervisor")
        .order_by(Lower("full_name"))
    )

    student_ids = [s.id for s in students]
    evidence_by_student: dict[int, dict[str, ProjectEvidence]] = {sid: {} for sid in student_ids}
    if student_ids:
        for ev in ProjectEvidence.objects.filter(student_id__in=student_ids):
            evidence_by_student[ev.student_id][ev.category] = ev

    pdf_bytes = build_class_evidence_pdf(
        school=profile.school,
        class_level=class_level,
        students=students,
        evidence_by_student=evidence_by_student,
    )
    filename = f"{profile.school.name}_{class_level}_evidence.pdf".replace(" ", "_")
    resp = HttpResponse(pdf_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _safe_folder_name(name: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", name, flags=re.UNICODE).strip()
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned[:80] or "student"


@login_required
def export_class_with_evidence(request, class_level: str):
    """
    ZIP download: Excel scores for the class + all student photo evidence files.
    """
    profile = _require_profile(request)
    if not _is_overall(profile):
        raise Http404()

    class_level = class_level.upper()
    if class_level not in {c for c, _ in SecondaryClassLevel.choices}:
        raise Http404()

    year = int((request.GET.get("year") or timezone.now().year))
    term = int((request.GET.get("term") or current_school_term()))
    scheme = AssessmentScheme.objects.filter(active=True).order_by("id").first()
    scheme_id = request.GET.get("scheme_id")
    if scheme_id and scheme_id.isdigit():
        scheme = AssessmentScheme.objects.filter(id=int(scheme_id)).first() or scheme
    if not scheme:
        raise Http404("No scheme configured")

    competencies = list(Competency.objects.filter(scheme=scheme).order_by("order", "id"))
    students = list(
        Student.objects.active_only().filter(school=profile.school, class_level=class_level)
        .select_related("supervisor")
        .order_by(Lower("full_name"))
    )
    student_ids = [s.id for s in students]

    evidence_by_student: dict[int, dict[str, ProjectEvidence]] = {sid: {} for sid in student_ids}
    if student_ids:
        for ev in ProjectEvidence.objects.filter(student_id__in=student_ids):
            evidence_by_student[ev.student_id][ev.category] = ev

    category_labels = [c.label for c in EvidenceCategory]
    wb = Workbook()
    ws = wb.active
    ws.title = f"{class_level} Scores"
    headers = (
        ["Student Name", "Supervisor", "Student No", "Stream"]
        + [c.name for c in competencies]
        + ["Evidence photos uploaded"]
        + category_labels
    )
    ws.append(headers)

    titles_wb = Workbook()
    titles_ws = titles_wb.active
    titles_ws.title = f"{class_level} Project Titles"
    titles_ws.append(["Student Name", "Project Title", "Supervisor", "Student No", "Stream"])

    checklist = resolve_checklist(profile.school_id, year=year, term=term, class_level=class_level)
    score_maps = student_display_score_maps(
        students,
        scheme,
        year,
        term,
        checklist_for=lambda _student: checklist,
    )
    for s in students:
        score_map = score_maps.get(s.id, {})
        by_cat = evidence_by_student.get(s.id, {})
        row = [
            s.full_name,
            s.supervisor.get_full_name() or s.supervisor.username,
            s.student_no,
            s.stream,
        ]
        row += [float(score_map[c.id]) if c.id in score_map else "" for c in competencies]
        row += [f"{len(by_cat)}/{len(EvidenceCategory)}"]
        row += ["Yes" if cat.value in by_cat else "No" for cat in EvidenceCategory]
        ws.append(row)
        titles_ws.append(
            [
                s.full_name,
                s.project_title,
                s.supervisor.get_full_name() or s.supervisor.username,
                s.student_no,
                s.stream,
            ]
        )

    excel_buf = io.BytesIO()
    wb.save(excel_buf)
    titles_buf = io.BytesIO()
    titles_wb.save(titles_buf)

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{class_level}_scores_and_evidence.xlsx", excel_buf.getvalue())
        zf.writestr(f"{class_level}_project_titles.xlsx", titles_buf.getvalue())

        for s in students:
            folder = f"evidence/{_safe_folder_name(s.full_name)}_{s.id}"
            for cat in EvidenceCategory:
                ev = evidence_by_student.get(s.id, {}).get(cat.value)
                if not ev or not ev.file:
                    continue
                original = ev.file.name.split("/")[-1]
                ext = original.rsplit(".", 1)[-1] if "." in original else "jpg"
                arcname = f"{folder}/{cat.value}.{ext}"
                try:
                    with ev.file.open("rb") as f:
                        zf.writestr(arcname, f.read())
                except OSError:
                    continue

    zip_buf.seek(0)
    base = f"{profile.school.name}_{class_level}_{year}_T{term}_with_evidence".replace(" ", "_")
    resp = HttpResponse(zip_buf.getvalue(), content_type="application/zip")
    resp["Content-Disposition"] = f'attachment; filename="{base}.zip"'
    return resp


@login_required
def school_password_change(request):
    profile = _require_profile(request)
    form = SchoolPasswordChangeForm(
        user=request.user, membership=profile, data=request.POST or None
    )

    if request.method == "POST" and form.is_valid():
        from .password_reset_utils import sync_account_password

        sync_account_password(
            request.user,
            form.cleaned_data["new_password1"],
            membership=profile,
        )
        update_session_auth_hash(request, request.user)
        messages.success(
            request,
            f"Password updated for {profile.school.name if profile.school_id else 'your account'}.",
        )
        return redirect("password_change_done")

    return render(
        request,
        "auth/password_change.html",
        {
            "form": form,
            "active_profile": profile,
            "must_change_password": profile.must_change_password,
        },
    )


@login_required
def school_password_change_done(request):
    _require_profile(request)
    return render(request, "auth/password_change_done.html")


@login_required
def account_settings(request):
    profile = _require_profile(request)
    username_form = UsernameChangeForm(user=request.user)

    if request.method == "POST":
        if "update_username" in request.POST:
            username_form = UsernameChangeForm(request.POST, user=request.user)
            if username_form.is_valid():
                username_form.save()
                messages.success(request, "Username updated.")
                return redirect("account_settings")
        messages.error(request, "Could not update username. Please fix the errors below.")

    return render(
        request,
        "auth/account_settings.html",
        {"username_form": username_form, "active_profile": profile},
    )


@login_required
def registration_pending(request):
    profile = _require_profile(request)
    if _is_superadmin(profile) or not profile.school_id:
        return redirect("home")
    if profile.school.status != SchoolRegistrationStatus.PENDING:
        return redirect("home")
    return render(request, "auth/registration_pending.html", {"school": profile.school})


def maintenance_notice(request):
    from .maintenance_utils import is_maintenance_mode, maintenance_message

    if not is_maintenance_mode():
        return redirect("home")
    return render(
        request,
        "auth/maintenance.html",
        {"maintenance_message": maintenance_message()},
    )


@login_required
def registration_deactivated(request):
    profile = _require_profile(request)
    if _is_superadmin(profile) or not profile.school_id:
        return redirect("home")
    if profile.school.status != SchoolRegistrationStatus.DEACTIVATED:
        return redirect("home")
    return render(
        request,
        "auth/registration_deactivated.html",
        {
            "school": profile.school,
            "permanent_deletion_at": permanent_deletion_at(profile.school),
            "grace_days": SCHOOL_DELETION_GRACE_DAYS,
        },
    )


@login_required
def registration_rejected(request):
    profile = _require_profile(request)
    if _is_superadmin(profile) or not profile.school_id:
        return redirect("home")
    if profile.school.status != SchoolRegistrationStatus.REJECTED:
        return redirect("home")
    return render(request, "auth/registration_rejected.html", {"school": profile.school})


@login_required
def superadmin_dashboard(request):
    profile = _require_profile(request)
    if not _is_superadmin(profile):
        raise Http404()

    purged = purge_expired_deactivated_schools()
    if purged:
        messages.info(
            request,
            f"Permanently deleted {len(purged)} school(s) that completed the "
            f"{SCHOOL_DELETION_GRACE_DAYS}-day deactivation period.",
        )

    status_filter = request.GET.get("status", "pending")
    schools = School.objects.select_related("overall_supervisor").order_by("-created_at")
    if status_filter in {s.value for s in SchoolRegistrationStatus}:
        schools = schools.filter(status=status_filter)

    school_rows = [
        (
            school,
            permanent_deletion_at(school)
            if school.status == SchoolRegistrationStatus.DEACTIVATED
            else None,
        )
        for school in schools
    ]

    counts = {
        "pending": School.objects.filter(status=SchoolRegistrationStatus.PENDING).count(),
        "approved": School.objects.filter(status=SchoolRegistrationStatus.APPROVED).count(),
        "rejected": School.objects.filter(status=SchoolRegistrationStatus.REJECTED).count(),
        "deactivated": School.objects.filter(status=SchoolRegistrationStatus.DEACTIVATED).count(),
    }

    from .maintenance_utils import (
        is_maintenance_mode,
        maintenance_forced_by_env,
        maintenance_message,
    )
    from .models import PlatformSettings

    platform_settings = PlatformSettings.get()

    return render(
        request,
        "superadmin/dashboard.html",
        {
            "school_rows": school_rows,
            "status_filter": status_filter,
            "counts": counts,
            "SchoolRegistrationStatus": SchoolRegistrationStatus,
            "grace_days": SCHOOL_DELETION_GRACE_DAYS,
            "maintenance_active": is_maintenance_mode(),
            "maintenance_db_enabled": platform_settings.maintenance_enabled,
            "maintenance_forced_by_env": maintenance_forced_by_env(),
            "maintenance_message_text": maintenance_message(),
            "maintenance_message_draft": platform_settings.maintenance_message,
        },
    )


@login_required
def superadmin_toggle_maintenance(request):
    profile = _require_profile(request)
    if not _is_superadmin(profile):
        raise Http404()
    if request.method != "POST":
        raise Http404()

    from .maintenance_utils import maintenance_forced_by_env
    from .models import PlatformSettings

    platform_settings = PlatformSettings.get()
    action = (request.POST.get("action") or "").strip().lower()
    message = (request.POST.get("maintenance_message") or "").strip()

    if action == "enable":
        platform_settings.maintenance_enabled = True
        if message:
            platform_settings.maintenance_message = message
        platform_settings.maintenance_updated_at = timezone.now()
        platform_settings.maintenance_updated_by = request.user
        platform_settings.save(
            update_fields=[
                "maintenance_enabled",
                "maintenance_message",
                "maintenance_updated_at",
                "maintenance_updated_by",
            ]
        )
        messages.success(
            request,
            "Maintenance mode is ON. Only platform superadmins can sign in until you turn it off.",
        )
    elif action == "disable":
        platform_settings.maintenance_enabled = False
        platform_settings.maintenance_updated_at = timezone.now()
        platform_settings.maintenance_updated_by = request.user
        platform_settings.save(
            update_fields=[
                "maintenance_enabled",
                "maintenance_updated_at",
                "maintenance_updated_by",
            ]
        )
        if maintenance_forced_by_env():
            messages.warning(
                request,
                "Dashboard maintenance is off, but MAINTENANCE_MODE is still set in server "
                "environment variables — users remain locked out until that is cleared.",
            )
        else:
            messages.success(request, "Maintenance mode is OFF. Users can sign in again.")
    else:
        messages.error(request, "Unknown maintenance action.")

    return redirect("superadmin_dashboard")


@login_required
def superadmin_school_detail(request, school_id: int):
    profile = _require_profile(request)
    if not _is_superadmin(profile):
        raise Http404()

    school = get_object_or_404(School.objects.select_related("overall_supervisor"), pk=school_id)
    reject_form = SchoolRejectForm()
    students_count = Student.objects.active_only().filter(school=school).count()
    supervisors_count = UserProfile.objects.filter(school=school, role=Role.SUPERVISOR).count()

    return render(
        request,
        "superadmin/school_detail.html",
        {
            "school": school,
            "reject_form": reject_form,
            "students_count": students_count,
            "supervisors_count": supervisors_count,
            "permanent_deletion_at": permanent_deletion_at(school),
            "grace_days": SCHOOL_DELETION_GRACE_DAYS,
        },
    )


@login_required
def superadmin_approve_school(request, school_id: int):
    profile = _require_profile(request)
    if not _is_superadmin(profile):
        raise Http404()
    if request.method != "POST":
        return redirect("superadmin_school_detail", school_id=school_id)

    school = get_object_or_404(School, pk=school_id)
    school.status = SchoolRegistrationStatus.APPROVED
    school.rejection_reason = ""
    school.reviewed_at = timezone.now()
    school.reviewed_by = request.user
    school.save(update_fields=["status", "rejection_reason", "reviewed_at", "reviewed_by"])
    messages.success(request, f'"{school.name}" has been approved.')
    return redirect("superadmin_dashboard")


@login_required
def superadmin_reject_school(request, school_id: int):
    profile = _require_profile(request)
    if not _is_superadmin(profile):
        raise Http404()
    if request.method != "POST":
        return redirect("superadmin_school_detail", school_id=school_id)

    school = get_object_or_404(School, pk=school_id)
    form = SchoolRejectForm(request.POST)
    reason = form.cleaned_data["rejection_reason"] if form.is_valid() else ""
    school.status = SchoolRegistrationStatus.REJECTED
    school.rejection_reason = reason
    school.reviewed_at = timezone.now()
    school.reviewed_by = request.user
    school.save(update_fields=["status", "rejection_reason", "reviewed_at", "reviewed_by"])
    messages.warning(request, f'"{school.name}" has been rejected.')
    return redirect("superadmin_dashboard")


@login_required
def superadmin_delete_school(request, school_id: int):
    profile = _require_profile(request)
    if not _is_superadmin(profile):
        raise Http404()

    school = get_object_or_404(School, pk=school_id)

    if request.method == "POST":
        form = SchoolDeleteConfirmForm(request.POST, expected_name=school.name)
        if form.is_valid():
            name = school.name
            deactivate_school(school=school, by_user=request.user)
            purge_date = permanent_deletion_at(school)
            messages.success(
                request,
                f'"{name}" has been deactivated. Users cannot sign in. '
                f"Permanent deletion is scheduled for {purge_date.strftime('%b %d, %Y') if purge_date else '30 days from now'}. "
                "You can reactivate the school before that date.",
            )
            return redirect("superadmin_school_detail", school_id=school_id)
    else:
        form = SchoolDeleteConfirmForm(expected_name=school.name)

    return render(
        request,
        "superadmin/delete_school_confirm.html",
        {
            "school": school,
            "form": form,
            "grace_days": SCHOOL_DELETION_GRACE_DAYS,
        },
    )


@login_required
def superadmin_reactivate_school(request, school_id: int):
    profile = _require_profile(request)
    if not _is_superadmin(profile):
        raise Http404()

    school = get_object_or_404(School, pk=school_id)
    if school.status != SchoolRegistrationStatus.DEACTIVATED:
        messages.warning(request, "Only deactivated schools can be reactivated.")
        return redirect("superadmin_school_detail", school_id=school_id)

    if request.method == "POST":
        restore_deactivated_school(school=school)
        messages.success(
            request,
            f'"{school.name}" has been reactivated. Users can sign in and use the system again.',
        )
        return redirect(f"{reverse('superadmin_dashboard')}?status=approved")

    return render(
        request,
        "superadmin/reactivate_school_confirm.html",
        {
            "school": school,
            "permanent_deletion_at": permanent_deletion_at(school),
            "grace_days": SCHOOL_DELETION_GRACE_DAYS,
        },
    )
