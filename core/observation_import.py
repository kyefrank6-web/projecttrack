from __future__ import annotations

import csv
import io
from typing import BinaryIO

from openpyxl import load_workbook

from .models import Competency, ObservationChecklist, ObservationCriterion
from .observation_defaults import CHECKBOX_MAX_RATING


EXPECTED_HEADERS = {
    "competency",
    "section_code",
    "section_title",
    "section_preamble",
    "observation",
}


def _normalize_header(h: str) -> str:
    return (h or "").strip().lower().replace(" ", "_")


def _row_dict(headers: list[str], values: list) -> dict[str, str]:
    out: dict[str, str] = {}
    for i, h in enumerate(headers):
        if not h:
            continue
        val = values[i] if i < len(values) else ""
        out[h] = ("" if val is None else str(val)).strip()
    return out


def import_structure_rows(checklist: ObservationChecklist, rows: list[dict[str, str]]) -> int:
    """Import observation rows; returns number of items created."""
    by_name = {
        c.name.lower(): c
        for c in Competency.objects.filter(scheme=checklist.scheme).order_by("order", "id")
    }
    ObservationCriterion.objects.filter(checklist=checklist).delete()
    order = 0
    created = 0
    for row in rows:
        comp_name = row.get("competency", "").strip()
        observation = row.get("observation", "").strip()
        if not comp_name or not observation:
            continue
        comp = by_name.get(comp_name.lower())
        if not comp:
            continue
        order += 1
        ObservationCriterion.objects.create(
            checklist=checklist,
            competency=comp,
            section_code=row.get("section_code", "").strip(),
            section_title=row.get("section_title", "").strip(),
            section_preamble=row.get("section_preamble", "").strip(),
            description=observation,
            order=order,
            max_rating=CHECKBOX_MAX_RATING,
        )
        created += 1
    return created


def import_structure_csv(checklist: ObservationChecklist, file_obj: BinaryIO) -> int:
    raw = file_obj.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return 0
    headers = [_normalize_header(h) for h in reader.fieldnames]
    rows: list[dict[str, str]] = []
    for r in reader:
        values = [r.get(orig, "") for orig in reader.fieldnames]
        d = _row_dict(headers, values)
        rows.append(
            {
                "competency": d.get("competency", ""),
                "section_code": d.get("section_code", ""),
                "section_title": d.get("section_title", ""),
                "section_preamble": d.get("section_preamble", ""),
                "observation": d.get("observation", d.get("observation_text", "")),
            }
        )
    return import_structure_rows(checklist, rows)


def import_structure_xlsx(checklist: ObservationChecklist, file_obj: BinaryIO) -> int:
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return 0
    headers = [_normalize_header(str(h or "")) for h in raw_rows[0]]
    rows: list[dict[str, str]] = []
    for r in raw_rows[1:]:
        cells = [("" if c is None else str(c)).strip() for c in r]
        d = _row_dict(headers, cells)
        rows.append(
            {
                "competency": d.get("competency", ""),
                "section_code": d.get("section_code", ""),
                "section_title": d.get("section_title", ""),
                "section_preamble": d.get("section_preamble", ""),
                "observation": d.get("observation", d.get("observation_text", "")),
            }
        )
    return import_structure_rows(checklist, rows)


def import_structure_file(checklist: ObservationChecklist, uploaded_file) -> int:
    name = (getattr(uploaded_file, "name", "") or "").lower()
    if name.endswith(".xlsx"):
        return import_structure_xlsx(checklist, uploaded_file)
    return import_structure_csv(checklist, uploaded_file)


def export_structure_xlsx_bytes(structure) -> bytes:
    """Build Excel bytes for a built-in or custom checklist structure (editable template)."""
    from io import BytesIO

    from openpyxl import Workbook

    from .observation_defaults import parse_observation_item

    wb = Workbook()
    ws = wb.active
    ws.title = "Observations"
    ws.append(["competency", "section_code", "section_title", "section_preamble", "observation"])
    for comp_name, code, title, preamble, observations in structure:
        for obs_item in observations:
            obs_text, max_rating = parse_observation_item(obs_item)
            if max_rating == 0:
                continue
            ws.append([comp_name, code, title, preamble, obs_text])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def builtin_structure_xlsx_filename(format_year: int) -> str:
    return f"UNEB_checklist_structure_{format_year}_format.xlsx"
